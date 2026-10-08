"""Normativas en Excel: conversión de las «clearance tables» de los ingenieros
(sin perder ningún texto), plantilla universal (exportar/importar ida y vuelta)
y validación de filas."""
import os

import openpyxl
import pytest

from nucleo import normativas as N
from nucleo import normativas_clearance as CL
from nucleo import normativas_excel as X

ORIGINAL = r"C:\Users\User\Downloads\Utility_Design_Aid_Clearance_Tables.xlsx"


def _libro_clearance(ruta):
    """Un libro chico con la misma forma que el de los ingenieros."""
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "E-2 ELECTRICAL"
    filas = [
        ["PROYECTO X — TYPICAL UTILITY CLEARANCES"],
        ["ELECTRICAL HORIZONTAL CLEARANCE TABLE"],
        ["WATER", "SEWER", "RAILROAD TRACK", "CURB & GUTTER"],
        ["5' MIN.¹", '18" MIN.²', "10' MIN. FROM CL OF TRACK³", "3' MIN. (LESS THAN 3' REQUIRES APPROVAL)¹"],
        [],
        ["ELECTRICAL VERTICAL CLEARANCE TABLE"],
        ["WATER", "RAILROAD TRACK", "COVER"],
        ["*1' MIN.¹", "4' BELOW TOP OF RAIL²", '36"-42"³'],
        ["*CONCRETE BLANKET REQUIRED WHEN LESS THAN 1'"],
        ["REFERENCES:"],
        ["1:  MWD GUIDELINES"], ["2:  STANDARD PLAN S-255-1"], ["3:  MRDC 4.1.4.5"],
        ["NOTES (GENERAL):"], ["1.  DISTRICT APPROVAL REQUIRED"], ["        · SUB NOTE"],
        ["PALABRA SUELTA"],
    ]
    for f in filas:
        ws.append(f)
    ws2 = wb.create_sheet("G-3 GAS VERT")
    for f in [["GAS VERTICAL CLEARANCE TABLE"], ["GAS MAIN TYPE", "WATER¹", "COVER"],
              ["HIGH PRESSURE MAIN LINE", '18" MIN.', "3' MIN.²"], [],
              ["REFERENCES:"], ["1:  MWD GUIDELINES"], ["2:  LA BOE NOTES"]]:
        ws2.append(f)
    wb.save(ruta)
    return ruta


@pytest.fixture
def clearance(tmp_path):
    return _libro_clearance(str(tmp_path / "clearance.xlsx"))


def test_conversion_interpreta_valores(clearance):
    reglas, anexos = CL.convertir(openpyxl.load_workbook(clearance))
    por = {(r["utilidades"][0], r["tipo"], r["params"]["contra"], r["params"]["condicion"]): r for r in reglas}
    agua = por[("ELECTRICO", "separacion_horizontal", "AGUA", "")]["params"]
    assert agua["minimo"] == 5 and agua["referencias"] == ["REF-01"] and agua["estado"] == "OK"
    assert por[("ELECTRICO", "separacion_horizontal", "ALCANTARILLADO", "")]["params"]["minimo"] == 1.5
    via = por[("ELECTRICO", "separacion_horizontal", "VIA_FERREA", "")]["params"]
    assert via["minimo"] == 10 and via["medido_desde"].startswith("Eje de vía")
    assert por[("ELECTRICO", "separacion_horizontal", "BORDILLO", "")]["params"]["estado"] == "Revisar"
    tope = por[("ELECTRICO", "separacion_vertical", "VIA_FERREA", "")]["params"]
    assert tope["minimo"] == 4 and tope["medido_desde"] == "Bajo el tope del riel"
    cub = por[("ELECTRICO", "recubrimiento", "SUPERFICIE", "")]["params"]
    assert (cub["minimo"], cub["maximo"]) == (3.0, 3.5)
    ast = por[("ELECTRICO", "separacion_vertical", "AGUA", "")]["params"]
    assert ast["notas"] and "CONCRETE BLANKET" in next(n["texto"] for n in anexos["notas"] if n["id"] == ast["notas"][0])
    gas = por[("GAS", "separacion_vertical", "AGUA", "Alta presión – línea principal")]["params"]
    assert gas["minimo"] == 1.5 and gas["referencias"] == ["REF-01"]       # superíndice del encabezado
    # la misma referencia en dos hojas = un solo id
    assert sum(1 for r in anexos["referencias"].values() if r["texto"] == "MWD GUIDELINES") == 1


def test_conversion_no_pierde_ningun_texto(clearance):
    wb = openpyxl.load_workbook(clearance)
    assert CL.celdas_no_usadas(wb) == []
    _reglas, anexos = CL.convertir(wb)
    textos = [n["texto"] for n in anexos["notas"]]
    assert any("PALABRA SUELTA" in t for t in textos) and any("SUB NOTE" in t for t in textos)
    assert anexos["documento"]["titulos"][0]["texto"].startswith("PROYECTO X")


def test_plantilla_ida_y_vuelta(clearance, tmp_path):
    rx, anexos, err, fmt = X.importar(clearance)
    assert fmt == "clearance" and not err
    reglas = N.cargar_catalogo(str(tmp_path / "no.json")); N.fusionar(reglas, rx)
    ruta = str(tmp_path / "plantilla.xlsx")
    X.exportar(reglas, anexos, ruta, "1.3.0")
    r2, a2, e2, f2 = X.importar(ruta)
    assert f2 == "plantilla" and e2 == [] and len(r2) == len(reglas)
    d2 = {r["id"]: r for r in r2}
    for r in reglas:
        q = d2[r["id"]]
        assert (q["tipo"], q["utilidades"], q["activa"], q["obligatoria"]) == \
            (r["tipo"], r["utilidades"], r["activa"], r["obligatoria"])
        for k, v in r["params"].items():
            if k != "descripcion":
                assert (q["params"].get(k) or None) == (v or None), (r["id"], k)
    assert a2["referencias"] == anexos["referencias"] and len(a2["notas"]) == len(anexos["notas"])


def test_validacion_de_filas(tmp_path):
    reglas = N.cargar_catalogo(str(tmp_path / "no.json"))
    ruta = str(tmp_path / "p.xlsx")
    X.exportar(reglas, {}, ruta)
    wb = openpyxl.load_workbook(ruta); ws = wb["Reglas"]
    col = {k: i + 1 for i, (k, *_r) in enumerate(X.COLUMNAS)}
    malas = [
        {"tipo": "Cualquier cosa", "utilidad": "Agua", "minimo": 1},
        {"tipo": "Separación horizontal", "utilidad": "Agua", "contra": "Gas", "minimo": -1},
        {"tipo": "Separación horizontal", "utilidad": "Agua", "contra": "Gas", "minimo": 5, "maximo": 2},
        {"tipo": "Separación horizontal", "utilidad": "Agua", "contra": "Gas"},
        {"tipo": "Ángulo de accesorio", "utilidad": "Gas", "contra": "Válvula", "angulos": "45"},
        {"tipo": "Separación horizontal", "utilidad": "Marte", "contra": "Gas", "minimo": 1},
        {"id": "conex_codo", "tipo": "Separación horizontal", "utilidad": "Gas", "contra": "Agua", "minimo": 1},
    ]
    buena = {"tipo": "Separación vertical", "utilidad": "Gas; Agua", "contra": "Vía férrea", "minimo": 6,
             "obligatoria": "No", "refs": "REF-99"}
    for f, datos in enumerate(malas + [buena], ws.max_row + 1):
        for k, v in datos.items():
            ws.cell(row=f, column=col[k], value=v)
    wb.save(ruta)
    r, _a, errores, _f = X.importar(ruta)
    assert len(errores) == len(malas)
    nueva = r[-1]
    assert nueva["utilidades"] == ["GAS", "AGUA"] and nueva["obligatoria"] is False
    assert nueva["params"]["estado"] == "Revisar"                  # cita una referencia que no existe


def test_importar_actualiza_las_de_la_app(tmp_path):
    reglas = N.cargar_catalogo(str(tmp_path / "no.json"))
    ruta = str(tmp_path / "p.xlsx")
    X.exportar(reglas, {}, ruta)
    wb = openpyxl.load_workbook(ruta); ws = wb["Reglas"]
    col = {k: i + 1 for i, (k, *_r) in enumerate(X.COLUMNAS)}
    ws.cell(row=2, column=col["angulos"], value="90; 135")         # conex_codo
    wb.save(ruta)
    r, _a, err, _f = X.importar(ruta)
    assert not err
    n_new, n_upd = N.fusionar(reglas, r)
    assert (n_new, n_upd) == (0, 4)
    assert next(x for x in reglas if x["id"] == "conex_codo")["params"]["angulos"] == [90.0, 135.0]


@pytest.mark.skipif(not os.path.isfile(ORIGINAL), reason="Excel de los ingenieros no disponible")
def test_excel_real_completo(tmp_path):
    wb = openpyxl.load_workbook(ORIGINAL, data_only=True)
    assert CL.celdas_no_usadas(wb) == []
    reglas, anexos = CL.convertir(wb)
    assert len(reglas) == 197 and len(anexos["referencias"]) == 28
    ruta = str(tmp_path / "p.xlsx")
    X.exportar(reglas, anexos, ruta)
    salida = " ¦ ".join(str(c.value) for ws in openpyxl.load_workbook(ruta) for row in ws.iter_rows()
                        for c in row if c.value not in (None, ""))
    import re
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                t = str(c.value or "").strip()
                if t and not t.upper().startswith("REFERENCES"):
                    assert re.sub(r"^\d+\s*:\s*", "", t) in salida, (ws.title, t)
