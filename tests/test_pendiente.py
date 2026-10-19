"""Pendiente (%) de cada utilidad: por tramo como Civil 3D (ficha de la Vista 3D), total en la
Tabla de datos y su Excel, y al reimportar: si la pendiente cambió y no coincide con las cotas,
no se importa nada."""
import os

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from nucleo import ficha3d, model_ops, tabla_datos, tabla_excel  # noqa: E402

FT = 0.5                                       # pies por px


def _p(**k):
    d = dict(layer="AGUA", pts=[(0, 0), (100, 0), (200, 0)], inv_start=-4.0, inv_end=-5.0)
    d.update(k)
    return d


def test_pendiente_por_tramo_y_total():
    p = _p()                                                  # 100 ft, de -4 a -5
    assert model_ops.pendiente_total(p, FT) == pytest.approx(-1.0)
    assert model_ops.pendientes_tramos(p, FT) == [pytest.approx(-1.0), pytest.approx(-1.0)]
    q = _p(vertex_inv_in={1: -4.2}, vertex_inv_out={1: -4.2})  # quiebre de pendiente en el vértice 1
    assert model_ops.pendientes_tramos(q, FT) == [pytest.approx(-0.4), pytest.approx(-1.6)]
    tb = tabla_datos.armar([q], [], [], [], lambda x, y: (x, y), FT)
    assert tb[0].filas[0][tb[0].columnas.index("Pendiente (%)")] == pytest.approx(-1.0)
    titulo, filas = ficha3d.de_utilidad(tb, [q], [], 0, FT)
    d = dict(filas)
    # con cotas que cambian por tramo, la de cada tramo EN LUGAR de la de punta a punta (pedido 2026-10-16)
    assert d["Pendiente por tramo (%)"] == "T1 -0.4 · T2 -1.6" and "Pendiente (%)" not in d
    igual = dict(ficha3d.de_utilidad(tb, [_p()], [], 0, FT)[1])                    # todos iguales: una sola
    assert "Pendiente por tramo (%)" not in igual and "Pendiente (%)" in igual


def _leido(**fila):
    return {"utilidades": [dict({"N°": 1}, **fila)]}


def test_pendiente_que_no_coincide_con_las_cotas_bloquea():
    pipes = [_p()]
    # sin tocar la pendiente (aunque cambie una rasante): se recalcula, no bloquea
    assert not tabla_excel.incoherencias(pipes, _leido(**{"Pendiente (%)": -1.0, "Rasante final (ft)": -6}), FT)
    # pendiente nueva SIN cambiar las cotas: bloquea y sugiere la rasante
    malas = tabla_excel.incoherencias(pipes, _leido(**{"Pendiente (%)": -2.0}), FT)
    assert len(malas) == 1 and "-6.000" in malas[0] and "-1.000" in malas[0]
    # pendiente nueva + rasante final que coincide: pasa (y solo se aplica la rasante)
    leido = _leido(**{"Pendiente (%)": "-2,0", "Rasante final (ft)": -6, "Rasante inicial (ft)": -4})
    assert not tabla_excel.incoherencias(pipes, leido, FT)
    cambios, errores = tabla_excel.cambios(pipes, [], leido)
    assert not errores and [(c.campo, c.despues) for c in cambios] == [("inv_end", -6.0)]
    assert tabla_excel.incoherencias(pipes, _leido(**{"Pendiente (%)": "mucha"}), FT)


@pytest.fixture
def win(monkeypatch):
    from PySide6 import QtWidgets
    QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    from traduccion import i18n, i18n_core
    monkeypatch.setattr(i18n, "_settings", lambda: type("S", (), {"setValue": lambda *a: None})())
    previo = i18n_core._current_lang
    i18n_core._current_lang = "es"
    from ui.ventana.app_window import Main
    w = Main(); w._info = lambda *a, **k: None
    ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)), "escenarios", "cruces", "escenarios_prueba.digproj")
    w._open_project_path(ruta); w._redraw()
    w.avisos = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "warning", lambda _p, _t, msg, *a: w.avisos.append(msg))
    yield w
    if getattr(w, "_tabla_dlg", None) is not None:
        w._tabla_dlg.close()
    w._dirty = False
    w.close()
    i18n_core._current_lang = previo


def test_excel_en_la_ventana_no_importa_si_la_pendiente_no_coincide(win, tmp_path):
    import openpyxl
    from ui.dialogos import tabla_datos_dialog
    dlg = tabla_datos_dialog.abrir(win)
    ruta = str(tmp_path / "datos.xlsx")
    tabla_excel.exportar(dlg.tablas, ruta)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Utilidades"]
    cab = [c.value for c in ws[1]]
    assert ws.cell(row=1, column=cab.index("Pendiente (%)") + 1).fill.fgColor.rgb.endswith(tabla_excel.VERDE)
    col_p, col_z = cab.index("Pendiente (%)") + 1, cab.index("Rasante final (ft)") + 1
    ws.cell(row=2, column=col_p).value = 3.25                           # pendiente nueva, cotas iguales
    ws.cell(row=3, column=col_z).value = -9.0                           # otra fila: cambio válido
    wb.save(ruta)
    antes = [dict(p) for p in win.pipes]
    assert dlg.importar(ruta, confirmar=False) == []                    # no se importa NADA
    assert win.avisos and "Utilidad #1" in win.avisos[-1]
    assert [p.get("inv_end") for p in win.pipes] == [p.get("inv_end") for p in antes]
    # corregida (la rasante final que da esa pendiente): se importa
    ft = win.scale / win.zoom
    p0 = win.pipes[0]
    ws.cell(row=2, column=col_z).value = round(float(p0["inv_start"] or 0) + 3.25 * model_ops.largo_planta_ft(p0, ft)
                                               / 100.0, 3)
    wb.save(ruta)
    aplicados = dlg.importar(ruta, confirmar=False)
    assert {(c.indice, c.campo) for c in aplicados} == {(0, "inv_end"), (1, "inv_end")}
    assert model_ops.pendiente_total(win.pipes[0], ft) == pytest.approx(3.25, abs=0.002)
