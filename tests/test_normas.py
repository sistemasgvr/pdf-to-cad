"""Normativas (tabla única de reglas desde 2026-10-21; antes cuatro tablas del Excel del usuario): catálogo,
Excel (ida y vuelta con desplegables; el Excel viejo de cuatro hojas se sigue importando), validación de
utilidades y accesorios, tabla de datos y la ventana real (Tipo + «+», Ver → Avisos, ventanas)."""
import os

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from nucleo import normas_catalogo as nc  # noqa: E402
from nucleo import normas_excel, normas_validar as nv  # noqa: E402
from nucleo import normas_excel_ayuda as ayuda, normas_reglas as NR  # noqa: E402

FT = 0.1


def _p(layer="AGUA", tipo="", diam=None, pts=((0, 0), (100, 0)), **k):
    d = dict(layer=layer, pts=[tuple(q) for q in pts], tipo=tipo, name="")
    if diam is not None:
        d.update(diam=float(diam), pipe_size=f"{diam:g} in")
    d.update(k)
    return d


def _reglas(cat, **k):
    return [r for r in cat["reglas"] if all(r.get(c) == v for c, v in k.items())]


# ─────────────────────────── catálogo y Excel ───────────────────────────

def test_lectura_de_celdas():
    assert nc.lista_numeros("6; 8; 12; 16") == [6, 8, 12, 16]
    assert nc.lista_numeros("1;1,5;2") == [1, 1.5, 2]
    assert nc.lista_numeros(4) == [4]
    assert nc.rango("0-320") == [0, 320] and nc.rango("101-200") == [101, 200] and nc.rango(None) is None
    assert nc.texto_lista([11.25, 22.5]) == "11,25; 22,5"
    assert nc.mismo("Instalación  domiciliaria", "INSTALACION DOMICILIARIA")


def test_tipos_y_agregar_tipo():
    cat = nc.nuevo()
    assert nc.tipos_de(cat, "AGUA")[:3] == ["DISTRIBUCION PRINCIPAL", "ALIMENTACION Y CONDUCCION",
                                             "INSTALACION DOMICILIARIA"]
    assert nc.tipos_de(cat, "TELECOM") == []
    assert nc.agregar_tipo(cat, "TELECOM", "Fibra  óptica") == "FIBRA OPTICA"
    nc.agregar_tipo(cat, "TELECOM", "fibra optica")                    # repetido: no se duplica
    assert nc.tipos_de(cat, "TELECOM") == ["FIBRA OPTICA"]


def test_guardar_y_cargar():
    cat = nc.nuevo()
    nc.agregar_tipo(cat, "TELECOM", "FIBRA")
    nc.guardar(cat)
    assert nc.tipos_de(nc.cargar(), "TELECOM") == ["FIBRA"]
    assert nc.cargar() == nc.normalizar(cat)


def test_de_fabrica_lleva_el_excel_del_ingeniero_y_los_word():
    """Excel del ingeniero (2026-10-21): manda; filas 8–11 = ALIMENTACION Y CONDUCCION. Lo que cambió se quitó."""
    cat = nc.nuevo()
    ac = "ALIMENTACION Y CONDUCCION"
    assert _reglas(cat, utilidad="AGUA", tipo=ac, elemento="Y", revisar="permitido")[0]["valores"] == ["NO"]
    assert _reglas(cat, utilidad="AGUA", tipo=ac, elemento="T", revisar="diam_principal")[0]["valores"] == [16, 20, 24]
    assert _reglas(cat, utilidad="AGUA", tipo=ac, elemento="CODO", revisar="angulo")[0]["valores"] == [11.25, 22.5, 45, 90]
    assert ac in _reglas(cat, utilidad="AGUA", elemento="T", revisar="solo_en")[0]["valores"]
    assert _reglas(cat, utilidad="AGUA", tipo="DISTRIBUCION PRINCIPAL", elemento="CRUZ",
                   revisar="diam_principal")[0]["valores"] == [6, 8, 12, 16]
    assert "DIMENSIONAMIENTO RESIDENCIAL" not in nc.tipos_de(cat, "ELECTRICO")      # renombrado en el Excel
    assert not _reglas(cat, revisar="radio", valores=[3, 4])                         # radio fijo: quitado
    amp = sorted(tuple(r["si_amp"]) for r in cat["reglas"] if r.get("si_amp"))
    assert amp == [(0, 320), (0, 320), (321, 400), (321, 400)]
    assert _reglas(cat, utilidad="GAS", tipo="DISTRIBUCION PRINCIPAL", revisar="diametro")[0]["min"] == 6
    # de los Word: separaciones agua–alcantarillado, buzones, pendientes y lo que no se mide como referencia
    assert _reglas(cat, utilidad="AGUA", revisar="sep_h", con_utilidad="ALCANTARILLADO")[0]["min"] == 10
    assert _reglas(cat, utilidad="AGUA", revisar="encima", con_utilidad="ALCANTARILLADO")[0]["min"] == 1
    assert {r["max"] for r in _reglas(cat, utilidad="ALCANTARILLADO", revisar="dist_buzones")} == {400, 500, 900}
    assert _reglas(cat, utilidad="DRENAJE", tipo="SUBDRENAJE", revisar="pendiente")[0]["min"] == 0.5
    assert all(r["fuente"] for r in cat["reglas"] if r["utilidad"] in ("ALCANTARILLADO", "DRENAJE"))
    assert _reglas(cat, revisar="referencia")


def test_catalogo_viejo_guardado_pasa_al_nuevo_con_los_tipos_del_usuario(tmp_path):
    import json
    ruta = tmp_path / "normativas_tablas.json"
    viejo = {"version": 1, "tipos": [["AGUA", "DISTRIBUCION PRINCIPAL"], ["AGUA", "RIEGO"],
                                     ["ELECTRICO", "DIMENSIONAMIENTO RESIDENCIAL"]],
             "diametros": [], "accesorios": [], "electricas": []}
    ruta.write_text(json.dumps(viejo), encoding="utf-8")
    cat = nc.cargar(str(ruta))
    assert "RIEGO" in nc.tipos_de(cat, "AGUA")                       # lo agregó el usuario: se queda
    assert "DIMENSIONAMIENTO RESIDENCIAL" not in nc.tipos_de(cat, "ELECTRICO")
    assert cat["reglas"] == nc.nuevo()["reglas"]
    assert (tmp_path / "normativas_tablas.v1.json").exists()          # respaldo del viejo
    assert json.loads(ruta.read_text(encoding="utf-8"))["version"] == 2


def test_excel_ida_y_vuelta_con_desplegables(tmp_path):
    import openpyxl
    cat = nc.nuevo()
    nc.agregar_tipo(cat, "ELECTRICO", "ALUMBRADO")
    ruta = tmp_path / "normas.xlsx"
    normas_excel.exportar(cat, ruta)
    problemas = []
    leido, faltan = normas_excel.importar(ruta, problemas)
    assert not faltan and not problemas and leido == nc.normalizar(cat)
    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == ["INSTRUCCIONES", "TIPOS", "REGLAS", "LISTAS"]
    assert wb["LISTAS"].sheet_state == "hidden" and "TablaReglas" in wb["REGLAS"].tables
    dvs = {str(dv.sqref): dv for dv in wb["REGLAS"].data_validations.dataValidation}
    assert "INDIRECT(\"E_\"&$B2)" in dvs["D2:D1000"].formula1                     # ELEMENTO según la utilidad
    assert "INDIRECT(\"R_\"&$B2" in dvs["L2:L1000"].formula1                      # REVISAR según los dos
    assert "OFFSET(TIPOS!$K$2" in dvs["C2:C1000"].formula1 and dvs["C2:C1000"].errorStyle == "stop"
    assert dvs["M2:M1000"].type == "decimal" and dvs["R2:R1000"].formula1 == '"ERROR,AVISO,INFO"'
    assert "R_AGUA_TDOMICILIARIA" in wb.defined_names and "E_ALCANTARILLADO" in wb.defined_names
    assert all(c.comment is not None for c in wb["REGLAS"][1])                    # cada columna explica qué va
    # INSTRUCCIONES explica cada columna y trae la tabla con TODO lo que se puede revisar
    texto = "\n".join(str(v) for fila in wb["INSTRUCCIONES"].iter_rows(values_only=True) for v in fila if v)
    for col in ("ACTIVA", "UTILIDAD", "TIPO", "ELEMENTO", "REVISAR", "MIN", "MAX", "VALORES", "CON UTILIDAD",
                "NIVEL", "FUENTE", "NOTA", "BUZON", "T - DOMICILIARIA", "CUALQUIERA"):
        assert col in texto, col
    assert all(r.texto in texto for r in NR.REVISAR)
    assert set(ayuda.EXPLICA) == {r.clave for r in NR.REVISAR}


def test_excel_viejo_de_cuatro_hojas_se_pasa_a_reglas(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRESION-DIAMETROS"
    ws.append(["UTILIDAD", "TIPO", "DIAMETRO (in)", "SOLO SE PERMITE EN"])
    ws.append(["AGUA", "RIEGO", "2; 3", None])
    ruta = tmp_path / "x.xlsx"
    wb.save(ruta)
    cat, faltan = normas_excel.importar(ruta)
    assert not faltan and nc.tipos_de(cat, "AGUA") == ["RIEGO"]       # el tipo de la tabla entra solo
    assert _reglas(cat, tipo="RIEGO", revisar="diametro")[0]["valores"] == [2, 3]


def test_excel_viejo_y_celdas_que_no_se_entienden(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "TIPOS"
    wb.active.append(["UTILIDAD", "TIPO"])
    wd = wb.create_sheet("PRESION-DIAMETROS")
    wd.append(["UTILIDAD", "TIPO", "DIAMETRO (in)", "SOLO SE PERMITE EN"])
    wd.append(["GAS", "SERVICIO RESIDENCIAL", "0.5-2", None])               # rango en una celda
    we = wb.create_sheet("ELECT-TELECOM")
    we.append(["UTILIDAD", "ESTRUCTURA", "TIPO", "AMPERAJE (A)", "LONGITUD RANGO (pies)", "DIAMETRO MIN (in)"])
    we.append(["ELECTRICO", "SIMPLE", "X", "0-320", "0-100", 3])            # formato viejo: se sigue leyendo
    we.append(["ELECTRICO", "SIMPLE", "X", ">=400", "0-100", 3])            # operador: no se entiende
    aux = wb.create_sheet("auxiliar")                                       # nunca se lee
    aux.append(["UTILIDAD", "TIPO"])
    aux.append(["AGUA", "IDEA"])
    ruta = tmp_path / "viejo.xlsx"
    wb.save(ruta)
    problemas = []
    cat, _f = normas_excel.importar(ruta, problemas)
    el = _reglas(cat, utilidad="ELECTRICO", revisar="diametro")
    assert el[0]["si_amp"] == [0, 320] and el[0]["si_largo"] == [0, 100] and el[1]["si_amp"] is None
    assert len(problemas) == 2
    assert any("PRESION-DIAMETROS, fila 2" in p and "0.5-2" in p for p in problemas)
    assert any("ELECT-TELECOM, fila 3" in p and ">=400" in p for p in problemas)
    assert "IDEA" not in nc.tipos_de(cat, "AGUA")


def test_excel_nuevo_rechaza_lo_que_la_app_no_contempla(tmp_path):
    import openpyxl
    ruta = str(tmp_path / "n.xlsx")
    normas_excel.exportar(nc.nuevo(), ruta)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["REGLAS"]
    n = ws.max_row
    ws.append(["SI", "AGUA", "TIPO INVENTADO", "TUBERIA", None, None, None, None, None, None, None, "Diámetro (in)",
               4])                                                         # tipo que no está en TIPOS
    ws.append(["SI", "AGUA", None, "BUZON", None, None, None, None, None, None, None, "Caída en el buzón, flujo "
               "directo (ft)", 0.1])                                       # buzones: solo en gravedad
    ws.append(["SI", "GAS", None, "TUBERIA", None, None, None, None, None, None, None, "Velocidad (fps)", 3])
    ws.append(["SI", "AGUA", None, "TUBERIA", None, None, None, None, None, None, None, "Material permitido",
               None, None, "Titanio"])                                     # material que la app no tiene
    wb["TIPOS"].append(["CABLE", "PRINCIPAL"])                              # utilidad que no es de la app
    wb.save(ruta)
    problemas = []
    cat, _f = normas_excel.importar(ruta, problemas)
    assert len(cat["reglas"]) == len(nc.nuevo()["reglas"])                # ninguna de las malas entra
    txt = "\n".join(problemas)
    assert f"REGLAS, fila {n + 1}, TIPO" in txt and f"REGLAS, fila {n + 2}, ELEMENTO" in txt
    assert f"REGLAS, fila {n + 3}, REVISAR" in txt and "Titanio" in txt and "CABLE" in txt


# ─────────────────────────── validación de utilidades ───────────────────────────

def _avisos(pipes, accs=(), cat=None):
    return nv.validar(cat or nc.nuevo(), pipes, list(accs), FT)


def test_sin_tipo_y_diametros_del_tipo():
    av = _avisos([_p(tipo=""), _p(tipo="DISTRIBUCION PRINCIPAL", diam=2), _p(tipo="DISTRIBUCION PRINCIPAL", diam=10),
                  _p(tipo="DISTRIBUCION PRINCIPAL", diam=8), _p(tipo="DISTRIBUCION PRINCIPAL", diam=4)])
    por = {}
    for a in av:
        por.setdefault(a.pipe, []).append(a)
    assert [a.clase for a in por[0]] == ["tipo"]
    assert "INSTALACION DOMICILIARIA" in por[1][0].mensaje                 # es de otro tipo
    assert "no permitido" in por[2][0].mensaje
    assert 3 not in por                                                     # 8" en principal: cumple
    assert por[4][0].info and "cul-de-sacs" in por[4][0].mensaje            # nota propia: informativo


def test_nota_que_repite_el_tipo_no_avisa():
    for t_, d in (("INSTALACION DOMICILIARIA", 1.5), ("ALIMENTACION Y CONDUCCION", 20),
                  ("DISTRIBUCION PRINCIPAL", 8)):
        assert _avisos([_p(tipo=t_, diam=d)]) == [], t_


def test_electrico_por_amperaje_y_largo():
    """Excel del ingeniero: instalación domiciliaria por amperaje Y largo (dos condiciones a la vez)."""
    dom, dp = "INSTALACION DOMICILIARIA", "DISTRIBUCION PRINCIPAL"
    largo = [(0, 0), (1500, 0)]                                             # 150 ft
    av = _avisos([_p("ELECTRICO", dom, 3, largo), _p("ELECTRICO", dom, 3, largo, amperaje=500),
                  _p("ELECTRICO", dom, 3, largo, amperaje=350), _p("ELECTRICO", dp, 1),
                  _p("TELECOM", "", 1), _p("TELECOM", "", 4), _p("ELECTRICO", dom, 4, largo, amperaje=350),
                  _p("ELECTRICO", dom, 3, largo, amperaje=320)])
    msgs = {a.pipe: a.mensaje for a in av if a.clase != "separacion"}     # (todas en el mismo lugar)
    assert "amperaje" in msgs[0].lower()
    assert 1 not in msgs                                                    # 500 A: ninguna fila
    assert 'debe ser 4"' in msgs[2] and "321-400 A" in msgs[2] and "101-200 ft" in msgs[2]
    assert 'mínimo 3"' in msgs[3] and 'mínimo 4"' in msgs[4]
    assert 5 not in msgs and 6 not in msgs and 7 not in msgs


def test_electrico_con_bancoducto_no_se_revisa():
    av = nv.validar(nc.nuevo(), [_p("TELECOM", "", 1)], [], FT, con_bancoducto={0})
    assert not av


def test_radio_de_curva_segun_el_diametro():
    """ESR Manual: conduit de 4 in o menos, radio mínimo 3 ft; mayor a 4 in, 4 ft."""
    pts = [(0, 0), (300, 0), (300, 300)]
    for d, radio, aviso in ((4, 3, False), (4, 2.5, True), (6, 3.5, True), (6, 4, False)):
        curva = [dict(x=300.0, y=0.0, curve=True, radius_ft=radio)]
        av = [a for a in nv.validar(nc.nuevo(), [_p("ELECTRICO", "DISTRIBUCION PRINCIPAL", d, pts)], [], FT,
                                    structures=curva) if a.clase == "radio"]
        assert bool(av) == aviso, (d, radio)
        if av:
            assert (av[0].x, av[0].y) == (300.0, 0.0) and "mínimo" in av[0].mensaje


# ─────────────────────────── accesorios ───────────────────────────

def _acc(tipo, capa="AGUA", tronco=(0,), ramal=None, angulo=90.0):
    return {"x": 0.0, "y": 0.0, "tipo": tipo, "angulo": angulo, "capa": capa, "red": capa,
            "pipes": sorted(set(tronco) | ({ramal} if ramal is not None else set())),
            "tronco": list(tronco), "ramal": ramal}


@pytest.mark.parametrize("d_tronco,d_ramal,esperado", [
    (12, 8, None), (12, 6, "debe ser 8"), (6, 4, "no hay un tamaño menor")])
def test_tee_ramal_un_tamano_menor(d_tronco, d_ramal, esperado):
    dp = "DISTRIBUCION PRINCIPAL"
    pipes = [_p(tipo=dp, diam=d_tronco), _p(tipo=dp, diam=d_ramal)]
    av = [a for a in _avisos(pipes, [_acc("tee", ramal=1)]) if a.clase == "accesorio"]
    if esperado is None:
        assert not av
    else:
        assert len(av) == 1 and esperado in av[0].mensaje


def test_tee_domiciliaria_solo_mira_el_ramal():
    pipes = [_p(tipo="DISTRIBUCION PRINCIPAL", diam=8), _p(tipo="INSTALACION DOMICILIARIA", diam=1),
             _p(tipo="INSTALACION DOMICILIARIA", diam=3)]
    assert not [a for a in _avisos(pipes, [_acc("tee", ramal=1)]) if a.clase == "accesorio"]
    av = [a for a in _avisos(pipes, [_acc("tee", ramal=2)]) if a.clase == "accesorio"]
    assert len(av) == 1 and "T domiciliaria" in av[0].mensaje
    av = [a for a in _avisos(pipes, [_acc("tee", ramal=1, angulo=30.0)]) if a.clase == "accesorio"]
    assert len(av) == 1 and "30°" in av[0].mensaje                          # ángulos 45; 90 del Excel


def test_wye_cruz_y_codos():
    dp, dom, ac = "DISTRIBUCION PRINCIPAL", "INSTALACION DOMICILIARIA", "ALIMENTACION Y CONDUCCION"
    pipes = [_p(tipo=dp, diam=10), _p(tipo=dom, diam=1), _p("GAS", "", 4), _p(tipo=ac, diam=20)]
    av = lambda a: [x.mensaje for x in _avisos(pipes, [a]) if x.clase == "accesorio"]  # noqa: E731
    assert "no permitida" in av(_acc("wye", ramal=0, tronco=(0,)))[0]
    assert "no permitida" in av(_acc("wye", ramal=0, tronco=(3,)))[0]       # fila nueva del Excel (8)
    assert '6", 8", 12", 16"' in av(_acc("cruz"))[0]                        # principal de 10": no
    assert not av(_acc("cruz", tronco=(3,)))                                # cruz en alimentación de 20": sí
    assert av(_acc("codo", angulo=30.0)) and not av(_acc("codo", angulo=45.2))
    assert av(_acc("codo", tronco=(3,), angulo=30.0))                       # fila nueva del Excel (11)
    assert not av(_acc("codo", tronco=(1,), angulo=30.0))                   # domiciliaria: cualquier ángulo
    assert av(_acc("codo", capa="GAS", tronco=(2,), angulo=30.0))
    assert "solo se permite en" in av(_acc("tee", tronco=(1,), ramal=None))[0]


def test_tee_real_desde_la_geometria():
    """Ramal que nace a mitad de la principal: `accesorios` dice cuál es el ramal."""
    from nucleo import accesorios
    dp = "DISTRIBUCION PRINCIPAL"
    pipes = [_p(tipo=dp, diam=12, pts=[(0, 0), (200, 0)]), _p(tipo=dp, diam=6, pts=[(100, 0), (100, 150)])]
    accs = accesorios.accesorios(pipes, None, FT)
    assert [(a["tipo"], a["ramal"], a["tronco"]) for a in accs] == [("tee", 1, [0])]
    av = [a for a in nv.validar(nc.nuevo(), pipes, accs, FT) if a.clase == "accesorio"]
    assert len(av) == 1 and "debe ser 8" in av[0].mensaje


# ─────────────────────────── rangos (catálogo de cuatro tablas pasado a reglas) ───────────────────────────

def test_rangos_min_max_y_abiertos():
    assert nc.rango_min_max("0,5", 2) == [0.5, 2] and nc.rango_min_max(400, None) == [400, None]
    assert nc.rango_min_max(None, "") is None and nc.rango_min_max(9, 3) == [3, 9]
    assert nc.en_rango(400, [400, None]) and nc.en_rango(10000, [400, None]) and not nc.en_rango(399, [400, None])
    assert nc.en_rango(5, [None, 5]) and not nc.en_rango(None, [0, 1]) and nc.en_rango(None, None)
    assert nc.texto_rango([400, None]) == "≥ 400" and nc.texto_rango([None, 100]) == "≤ 100"
    assert nc.texto_rango([0, 320]) == "0-320"


def _cat_rangos():
    viejo = {"tipos": [], "accesorios": [],
             "diametros": [{"utilidad": "GAS", "tipo": "SERVICIO RESIDENCIAL", "diametros": [], "diam_min": 0.5,
                            "diam_max": 2, "nota": ""},
                           {"utilidad": "GAS", "tipo": "DISTRIBUCION PRINCIPAL", "diametros": [], "diam_min": 6,
                            "diam_max": None, "nota": ""}],
             "electricas": [{"utilidad": "ELECTRICO", "estructura": "SIMPLE", "tipo": "DIMENSIONAMIENTO RESIDENCIAL",
                             "amperaje": [400, None], "longitud": [0, 100], "diam_fijo": [], "diam_min": 4,
                             "diam_max": None, "radio_min": 3, "radio_max": 4, "radio_fijo": []}]}
    return nc.normalizar(viejo)                                           # (se pasa solo a reglas)


def test_excel_con_rangos_ida_y_vuelta(tmp_path):
    cat = _cat_rangos()
    ruta = tmp_path / "n.xlsx"
    normas_excel.exportar(cat, ruta)
    problemas = []
    leido, _f = normas_excel.importar(ruta, problemas)
    assert not problemas and leido == cat


def test_validar_rangos_de_diametro_amperaje_y_radio():
    cat = _cat_rangos()
    gas = [_p("GAS", "SERVICIO RESIDENCIAL", 1.5), _p("GAS", "SERVICIO RESIDENCIAL", 3),
           _p("GAS", "DISTRIBUCION PRINCIPAL", 8), _p("GAS", "DISTRIBUCION PRINCIPAL", 4)]
    av = nv.validar(cat, gas, [], FT)
    assert [a.pipe for a in av if a.clase == "diametro"] == [1, 3]
    assert 'entre 0.5" y 2"' in next(a.mensaje for a in av if a.pipe == 1)
    assert 'mínimo 6"' in next(a.mensaje for a in av if a.pipe == 3)
    dom = "DIMENSIONAMIENTO RESIDENCIAL"
    pts = ((0, 0), (300, 0), (300, 300))
    elec = [_p("ELECTRICO", dom, 3, pts, amperaje=500), _p("ELECTRICO", dom, 3, pts, amperaje=350)]
    curvas = [dict(x=300.0, y=0.0, curve=True, radius_ft=10.0)]
    av = nv.validar(cat, elec, [], FT, structures=curvas)
    assert sorted((a.pipe, a.clase) for a in av) == [(0, "electrico"), (0, "radio")]
    radio = next(a for a in av if a.clase == "radio")
    assert (radio.x, radio.y) == (300.0, 0.0) and "entre 3 ft y 4 ft" in radio.mensaje
    curvas[0]["radius_ft"] = 3.5
    assert not [a for a in nv.validar(cat, elec, [], FT, structures=curvas) if a.clase == "radio"]


# ─────────────────────────── tabla de datos ───────────────────────────

def test_tabla_de_datos_y_excel(tmp_path):
    import openpyxl
    from nucleo import tabla_datos
    from nucleo.duct_bank import DuctBank
    pipes = [_p(tipo="DISTRIBUCION PRINCIPAL", diam=8), _p("ELECTRICO", "", pts=[(0, 50), (300, 50)])]
    structs = [{"cod": "BZ-1", "x": 0.0, "y": 0.0, "net": "gravity", "rim": 100.0, "sump": 95.0},
               {"cod": "CV-1", "x": 300.0, "y": 50.0, "curve": True, "radius_ft": 5.0}]
    db = DuctBank(name="DB1"); db.assign([1])
    av = nv.validar(nc.nuevo(), pipes, [], FT)
    tablas = tabla_datos.armar(pipes, structs, [db], av, lambda x, y: (x * FT, -y * FT), FT)
    ut, bz, cu, bd, avs, co = tablas
    assert [f[0] for f in ut.filas] == [1, 2] and ut.filas[0][3] == "DISTRIBUCION PRINCIPAL"
    assert ut.filas[0][9] == pytest.approx(10.0) and ut.filas[1][16] == "DB1"
    assert [f[0] for f in bz.filas] == ["BZ-1"] and len(cu.filas) == 1 and len(bd.filas) == 1
    assert len(avs.filas) == 1                                               # la eléctrica sin tipo
    ruta = tmp_path / "datos.xlsx"
    tabla_datos.exportar_excel(tablas, ruta)
    wb = openpyxl.load_workbook(ruta)
    assert len(wb.sheetnames) == 7 and wb.worksheets[1].max_row == 3      # «Léeme» + 6 tablas


# ─────────────────────────── ventana ───────────────────────────

@pytest.fixture
def win(monkeypatch):
    from PySide6 import QtGui, QtWidgets
    from traduccion import i18n, i18n_core
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])  # noqa: F841
    monkeypatch.setattr(i18n, "_settings", lambda: type("S", (), {"setValue": lambda *a: None})())
    previo = i18n_core._current_lang
    i18n_core._current_lang = "es"
    from ui.ventana.app_window import Main
    w = Main()
    img = QtGui.QImage(1500, 1500, QtGui.QImage.Format_RGB32); img.fill(QtGui.QColor(255, 255, 255))
    w.canvas.set_image(img)
    w.scale, w.zoom = 0.35, 3.5
    w.pipes = [_p("AGUA", "", 8, pts=[(100, 100), (600, 100)]), _p("ELECTRICO", "", pts=[(100, 300), (600, 300)])]
    w._refresh_lists(); w._redraw()
    w.avisos = []
    monkeypatch.setattr(QtWidgets.QMessageBox, "information", lambda _p, _t, msg, *a: w.avisos.append(msg))
    yield w
    w._dirty = False
    w.close()
    i18n_core._current_lang = previo


def test_ventana_tipo_avisos_y_mas(win, monkeypatch):
    from PySide6 import QtWidgets
    assert sum(1 for a in win._normas_avisos if a.clase == "tipo") == 2
    assert win.btn_normas.isVisibleTo(win) or win.btn_normas.text()
    win._show_tab(0); win.pipe_list.setCurrentRow(0)
    assert [win.prop_tipo.itemText(k) for k in range(win.prop_tipo.count())][1:] == nc.tipos_de(nc.nuevo(), "AGUA")
    assert not win.prop_amp.isVisibleTo(win.gprop)
    win.prop_tipo.setCurrentIndex(win.prop_tipo.findData("DISTRIBUCION PRINCIPAL"))
    assert win.pipes[0]["tipo"] == "DISTRIBUCION PRINCIPAL"
    assert not [a for a in win._normas_avisos if a.pipe == 0]
    monkeypatch.setattr(QtWidgets.QInputDialog, "getText", lambda *a, **k: ("Riego", True))
    win._agregar_tipo()
    assert win.pipes[0]["tipo"] == "RIEGO" and "RIEGO" in nc.tipos_de(nc.cargar(), "AGUA")
    win.undo()
    assert win.pipes[0]["tipo"] == "DISTRIBUCION PRINCIPAL"


def test_ver_avisos_y_ventanas(win):
    win._mostrar_avisos(False)
    assert not (win.chk_show_conflicts.isChecked() or win.act_show_acc.isChecked() or win.act_show_normas.isChecked())
    win._mostrar_avisos(True)
    assert [a.text() for a in win.menu_avisos.actions() if a.text()][-2:] == ["Mostrar todos", "Ocultar todos"]
    win.open_normativas(en_avisos=True)
    assert win._normas_dlg.lst_avisos.topLevelItemCount() == 2
    win.abrir_tabla_datos()
    assert win._tabla_dlg.pestanas.count() == 7          # + conflictos y cotas por vértice
    win._tabla_dlg.close(); win._normas_dlg.close()
