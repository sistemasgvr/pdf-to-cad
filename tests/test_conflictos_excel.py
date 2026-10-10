"""Conflictos numerados (C-n) en el lienzo, la Tabla de datos y la Vista 3D, y el Excel
de la Tabla de datos de IDA Y VUELTA (editar cotas en Excel y volver a cargarlas)."""
import os

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from nucleo import conflictos  # noqa: E402
from nucleo import tabla_excel  # noqa: E402

ESC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "escenarios", "escenarios_prueba.digproj")


def test_numerar_en_orden_de_lectura():
    items = [dict(clave="choque", x=500, y=100), dict(clave="cruce", x=10, y=400), dict(clave="exceso", x=20, y=102)]
    cs = conflictos.numerar(items)
    assert [(c.id, c.clave) for c in cs] == [("C-1", "exceso"), ("C-2", "choque"), ("C-3", "cruce")]
    # el mismo cruce visto por varios pares de tramos (en un vértice) es UN conflicto
    dup = [dict(clave="cruce", x=0, y=0, pipes=[1, 2]), dict(clave="cruce", x=0.4, y=0, pipes=[2, 1]),
           dict(clave="cruce", x=0, y=0, pipes=[1, 3])]
    assert len(conflictos.numerar(dup)) == 2
    c = conflictos.numerar([dict(clave="cruce", x=0, y=0, pipes=[1, 2], cotas=[-4.0, -4.05])])[0]
    assert c.diferencia == pytest.approx(0.05) and c.titulo == "Cruce a la misma cota"


@pytest.fixture
def win(monkeypatch):
    from PySide6 import QtWidgets
    QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    from traduccion import i18n, i18n_core
    monkeypatch.setattr(i18n, "_settings", lambda: type("S", (), {"setValue": lambda *a: None})())
    previo = i18n_core._current_lang
    i18n_core._current_lang = "es"
    from ui.ventana.app_window import Main
    w = Main(); w._info = lambda *a, **k: None
    w._open_project_path(ESC)
    w._redraw()
    yield w
    for d in ("_tabla_dlg", "_vista3d_dlg"):
        if getattr(w, d, None) is not None:
            getattr(w, d).close()
    w._dirty = False
    w.close()
    i18n_core._current_lang = previo


def test_conflictos_numerados_en_el_lienzo_y_la_tabla(win):
    cs = win._conflictos
    assert len(cs) >= 10 and [c.id for c in cs] == [f"C-{n}" for n in range(1, len(cs) + 1)]
    assert {c.clave for c in cs} >= {"cruce", "exceso", "escalon", "redes", "choque"}
    etiquetas = [it.data(0) for it in win._overlay if hasattr(it, "data") and str(it.data(0) or "").startswith("C-")]
    assert sorted(i for e in etiquetas for i in e.split()) == sorted(c.id for c in cs)  # cada uno, con su número en 2D
    tb = next(t for t in win._tablas_de_datos() if t.clave == "conflictos")
    assert [f[0] for f in tb.filas] == [c.id for c in cs]


def test_excel_de_ida_y_vuelta(win, tmp_path):
    import openpyxl
    from ui.dialogos import tabla_datos_dialog
    dlg = tabla_datos_dialog.abrir(win)
    ruta = str(tmp_path / "datos.xlsx")
    tabla_excel.exportar(dlg.tablas, ruta)
    assert tabla_excel.cambios(win.pipes, win.structures, tabla_excel.importar(ruta)) == ([], [])  # nada cambió
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Utilidades"]
    cab = [c.value for c in ws[1]]
    assert ws.cell(row=1, column=cab.index("Rasante inicial (ft)") + 1).fill.fgColor.rgb.endswith(tabla_excel.VERDE)
    ws.cell(row=2, column=cab.index("Rasante inicial (ft)") + 1).value = -12.5          # utilidad #1
    ws.cell(row=2, column=cab.index("Tipo") + 1).value = "DISTRIBUCION PRINCIPAL"
    wv = wb["Cotas por vértice"]
    cv = [c.value for c in wv[1]]
    fila = next(r for r in range(2, wv.max_row + 1))
    n_ut, vert = wv.cell(row=fila, column=1).value, wv.cell(row=fila, column=cv.index("Vértice") + 1).value
    wv.cell(row=fila, column=cv.index("Cota de salida (ft)") + 1).value = "-7,25"          # coma decimal
    wb.save(ruta)
    cambios, errores = tabla_excel.cambios(win.pipes, win.structures, tabla_excel.importar(ruta))
    assert not errores
    assert {(c.objeto, c.indice, c.campo) for c in cambios} == {
        ("pipe", 0, "inv_start"), ("pipe", 0, "tipo"), ("vertice", n_ut - 1, "vertex_inv_out")}
    aplicados = dlg.importar(ruta, confirmar=False)                          # como la ventana: con deshacer
    assert len(aplicados) == 3
    assert win.pipes[0]["inv_start"] == -12.5 and win.pipes[0]["tipo"] == "DISTRIBUCION PRINCIPAL"
    assert win.pipes[n_ut - 1]["vertex_inv_out"][vert] == pytest.approx(-7.25)
    win.undo()
    assert win.pipes[0]["inv_start"] != -12.5


def test_excel_con_errores_no_aplica_lo_malo(win, tmp_path):
    import openpyxl
    ruta = str(tmp_path / "datos.xlsx")
    from ui.dialogos import tabla_datos_dialog
    tabla_excel.exportar(tabla_datos_dialog.abrir(win).tablas, ruta)
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Utilidades"]
    cab = [c.value for c in ws[1]]
    ws.cell(row=2, column=cab.index("Rasante final (ft)") + 1).value = "abc"
    ws.cell(row=3, column=1).value = 999                                    # utilidad que no existe
    wb.save(ruta)
    cambios, errores = tabla_excel.cambios(win.pipes, win.structures, tabla_excel.importar(ruta))
    assert not cambios and len(errores) == 2


def test_conflictos_en_la_vista_3d(win):
    import time
    from PySide6 import QtWidgets
    from nucleo import ficha3d, modelo3d
    win._write_project = lambda *a, **k: None
    dlg = win.abrir_vista3d()
    t0 = time.time()
    while (dlg._trabajo is not None or dlg._t.isActive()) and time.time() - t0 < 20:
        QtWidgets.QApplication.processEvents(); time.sleep(0.01)
    assert dlg.visor.toolTip() == ""                                         # sin el globo de instrucciones
    objs = [(k, o) for k, o in dlg.escena.objetos.items() if o["tipo"] == "conflicto"]
    assert len(objs) == len(win._conflictos) >= len(dlg.escena.etiquetas) > 0
    assert {i for _i, _p, txt in dlg.escena.etiquetas for i in txt.split(" · ")} == {c.id for c in win._conflictos}
    k, o = objs[0]
    assert modelo3d.objeto_de(k) == ("conflicto", o["indice"])
    titulo, filas = ficha3d.de_conflicto(o)
    assert titulo.startswith(o["codigo"]) and dict(filas)["Conflicto"] == o["codigo"]
    dlg._elegido(k)                                                          # clic: ficha del conflicto
    assert dlg.ficha.lbl_titulo.text().startswith(o["codigo"])
    assert dlg.texto_de(k).startswith(o["codigo"])
