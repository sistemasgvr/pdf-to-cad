"""Formato SIMPLE de normativas (el de los ingenieros: «All (Flat)» + «References»)."""
import openpyxl
import pytest

from nucleo import normativas_excel as E
from nucleo import normativas_simple as S

CAB = ["Utility A", "Utility B", "Orientation", "Case / Sub-type", "Min (value)", "Min (unit)",
       "Max (value)", "Max (unit)", "Measured From", "Notes", "Reference(s)", "Source Sheet"]
FILAS = [
    ("WATER", "SEWER", "Horizontal", "General", 10, "ft", None, None, None, "Needs approval", "W-2:1", "W-2"),
    ("WATER", "COVER", "Vertical", 'Water main < 12"', 36, "in", 42, "in", None, None, "W-2:9", "W-2"),
    ("ELECTRICAL", "RAILROAD TRACK", "Vertical", "General", 4, "ft", None, None, "BELOW TOP OF RAIL", None, "E-2:7", "E-2"),
    ("SEWER", "RAILROAD TRACK", "Horizontal", "Perpendicular pipe crossing", None, None, None, None, None,
     "Concrete encasement 10' beyond track centerline.", "S-2:10", "S-2"),
    ("TELECOM", "BIKE LANE", "Horizontal", "General", 2, "ft", None, None, None, None, "T-2:1", "T-2"),
]


@pytest.fixture
def libro(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Legend"
    ws = wb.create_sheet("All (Flat)")
    ws.append(CAB)
    for f in FILAS:
        ws.append(f)
    wr = wb.create_sheet("References")
    wr.append(["Source Sheet", "Ref #", "Standard / Note Text"])
    for h, n in (("W-2", "1"), ("W-2", "9"), ("E-2", "7"), ("S-2", "10")):
        wr.append([h, n, f"Texto {h} {n}"])
    ruta = tmp_path / "simple.xlsx"
    wb.save(ruta)
    return str(ruta)


def test_importa_formato_simple(libro):
    reglas, anexos, errores, formato = E.importar(libro)
    assert formato == "simple" and errores == []
    por = {(r["utilidades"][0], r["params"]["contra"], r["tipo"]): r for r in reglas}
    agua = por[("AGUA", "ALCANTARILLADO", "separacion_horizontal")]["params"]
    assert agua["minimo"] == 10 and agua["referencias"] == ["W-2:1"] and agua["estado"] == "OK"
    cub = por[("AGUA", "SUPERFICIE", "recubrimiento")]["params"]
    assert cub["minimo"] == 3 and cub["maximo"] == 3.5 and cub["condicion"] == 'Water main < 12"'
    assert por[("ELECTRICO", "VIA_FERREA", "separacion_vertical")]["params"]["medido_desde"] == "BELOW TOP OF RAIL"
    assert por[("ALCANTARILLADO", "VIA_FERREA", "requisito")]["params"]["descripcion"].startswith("Concrete")
    raro = por[("TELECOM", "BIKE LANE", "separacion_horizontal")]["params"]
    assert raro["estado"] == "Revisar"                        # elemento desconocido + referencia que falta
    assert anexos["referencias"]["W-2:9"]["texto"] == "Texto W-2 9"


def test_ida_y_vuelta_no_pierde_nada(libro, tmp_path):
    reglas, anexos, _e, _f = E.importar(libro)
    reglas.append({"id": "ANG-1", "tipo": "angulos_accesorio", "utilidades": ["AGUA"], "activa": False,
                   "obligatoria": True, "params": {"accesorio": "codo", "angulos": [90.0, 135.0], "tolerancia": 1.0,
                                                   "descripcion": ""}})
    salida = tmp_path / "rt.xlsx"
    S.exportar(reglas, anexos, str(salida), "1.3.0")
    wb = openpyxl.load_workbook(salida)
    assert {"Legend", "All (Flat)", "References", "WATER", "Fitting Angles"} <= set(wb.sheetnames)
    fila = [c.value for c in wb["All (Flat)"][1]][:12]
    assert fila == CAB
    r2, a2, e2, f2 = E.importar(str(salida))
    assert f2 == "simple" and e2 == []
    clave = lambda r: (r["id"], r["tipo"], r.get("activa"), repr(sorted((r.get("params") or {}).items())))
    assert sorted(map(clave, r2)) == sorted(map(clave, reglas))
    assert a2["referencias"] == anexos["referencias"]


def test_errores_por_fila(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All (Flat)"
    ws.append(CAB)
    ws.append(("PIZZA", "WATER", "Horizontal", "General", 1, "ft"))
    ws.append(("WATER", "GAS", "Diagonal", "General", 1, "ft"))
    ws.append(("WATER", "GAS", "Horizontal", "General", 1, "yd"))
    ws.append(("WATER", "GAS", "Horizontal", "General", 5, "ft", 3, "ft"))
    ruta = tmp_path / "mal.xlsx"
    wb.save(ruta)
    reglas, _a, errores, _f = E.importar(str(ruta))
    assert reglas == [] and [f for f, _m in errores] == [2, 3, 4, 5]


def test_unidades_al_exportar():
    assert S._en_unidad([1.5, None]) == ([18, None], "in")
    assert S._en_unidad([3.0, 3.5]) == ([36, 42], "in")
    assert S._en_unidad([13.58, None]) == ([13.58, None], "ft")
    assert S._en_unidad([1.0, None], "in") == ([12, None], "in")
