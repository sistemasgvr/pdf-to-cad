"""IA (pedido del usuario 2026-10-23): «Pregúntale al proyecto», el informe redactado por la IA y lo que aprende
en una carpeta de texto plano junto al proyecto. Con un servidor Ollama SIMULADO (`ollama_falso`): sin red."""
import json
import os
import shutil
import time

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from ia import consulta, contexto, memoria as M, proveedor as P, redactor_ia, texto as T  # noqa: E402
from nucleo import informe_datos  # noqa: E402
from nucleo.conflictos import Conflicto  # noqa: E402
from nucleo.informe_datos import Opciones  # noqa: E402
from nucleo.normas_validar import Aviso  # noqa: E402
from ollama_falso import OllamaFalso  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
ESC = os.path.join(AQUI, "escenarios", "cruces", "escenarios_prueba.digproj")
FT = 0.1


def _pipes():
    return [{"layer": "AGUA", "pts": [(0, 0), (5000, 0)], "diam": 8, "tipo": "DISTRIBUCION PRINCIPAL",
             "material": "PVC", "inv_start": 100.0, "inv_end": 99.0},
            {"layer": "ALCANTARILLADO", "pts": [(100, -50), (100, 50)], "diam": 12},
            {"layer": "DRENAJE", "pts": [(4000, 0), (4000, 900)]}]


def _informe():
    conf = [Conflicto("C-1", "cruce", 100.0, 0.0, "Cruce sin conexión.", [0, 1], [101.0, 100.5])]
    av = [Aviso(4010.0, 400.0, "diametro", "Diámetro 6\" no permitido", pipe=0, fuente="Manual de agua")]
    return informe_datos.armar(conf, av, _pipes(), FT, Opciones())


def _datos(inf=None):
    inf = inf or _informe()
    return contexto.DatosProyecto.armar(_pipes(), inf.fichas, FT, "Prueba")


def _cfg(srv, **k):
    return P.Config.de({"proveedor": "ollama", "url": srv.url, **k})


# ── texto ──
def test_cifras_codigos_y_json():
    assert T.numeros("AGUA-5 cruza C-12 a 0,4 ft y 101.25 ft") == [0.4, 101.25]
    assert T.no_respaldados("baja 0.40 ft, sube 3.7 ft y las 2 tuberías", [0.4]) == [3.7]
    assert T.leer_json('```json\n{"a": 1}\n```') == {"a": 1} and T.leer_json("nada") is None
    assert "tuberia" in T.raices("Tuberías") or "tuber" in "".join(T.raices("Tuberías"))


# ── memoria (texto plano junto al proyecto) ──
def test_memoria_sin_proyecto_se_escribe_al_guardar_y_sigue_al_guardar_como(tmp_path):
    mem = M.Memoria()
    mem.ensenar(M.NOTA, "Los cruces con alcantarillado se resuelven bajando el agua")
    mem.ensenar(M.NOTA, "El gas va siempre por la vereda norte")
    assert mem.carpeta is None and not list(tmp_path.iterdir())
    carpeta = M.carpeta_de(str(tmp_path / "Obra.digproj"))
    assert carpeta.endswith("Obra - IA")
    mem.vincular(carpeta)
    lineas = open(os.path.join(carpeta, M.ENSENANZAS), encoding="utf-8").read().splitlines()
    assert len(lineas) == 2 and json.loads(lineas[0])["tipo"] == "nota"
    assert os.path.isfile(os.path.join(carpeta, "LEEME.txt"))
    assert mem.buscar("¿cómo resuelvo el cruce con el alcantarillado?")[0]["texto"].startswith("Los cruces")
    # una línea rota (editada a mano) no tumba la memoria
    with open(os.path.join(carpeta, M.ENSENANZAS), "a", encoding="utf-8") as f:
        f.write("{esto no es json\n")
    assert len(M.Memoria(carpeta).registros(M.ENSENANZAS)) == 2
    nueva = M.carpeta_de(str(tmp_path / "Obra v2.digproj"))
    mem.copiar_a(nueva)
    mem.ensenar(M.NOTA, "solo en la v2")
    assert len(M.Memoria(nueva).notas()) == 3 and len(M.Memoria(carpeta).notas()) == 2


def test_valoradas_como_no_utiles_no_vuelven():
    mem = M.Memoria()
    mem.ensenar(M.VALORACION, "", valor=-1, pregunta="cruce agua", respuesta="mal")
    mem.ensenar(M.VALORACION, "", valor=1, pregunta="cruce agua drenaje", respuesta="bien")
    assert [r["respuesta"] for r in mem.buscar("cruce de agua")] == ["bien"]


# ── contexto ──
def test_contexto_resume_y_elige_lo_que_viene_al_caso():
    d = _datos()
    assert "Agua (AGUA): 1 utilidades" in d.resumen and "Conflictos (C-n): 1" in d.resumen
    f, u = d.seleccionar("¿qué pasa en C-1?")
    assert f[0].codigo == "C-1" and {x.codigo for x in u} == {"AGUA-1", "ALCANTARILLADO-2"}
    f, u = d.seleccionar("¿cómo está el drenaje?")
    assert [x.codigo for x in u] == ["DRENAJE-3"] and not f
    f, _u = d.seleccionar("zona 2")
    assert [x.codigo for x in f] == ["N-1"]
    assert d.ref("agua-1")["objeto"] == "pipe" and d.ref("C-1")["x"] == 100.0


# ── proveedor ──
def test_ollama_trae_el_texto_en_trozos_y_avisa_errores():
    with OllamaFalso(lambda c: "Respuesta larga de prueba") as srv:
        prov = P.crear(_cfg(srv))
        trozos = []
        r = prov.chat([{"role": "user", "content": "hola"}], sistema="s", json=True, al_texto=trozos.append)
        assert r.texto == "Respuesta larga de prueba" and len(trozos) > 1 and "".join(trozos) == r.texto
        cuerpo = srv.pedidos[-1]
        assert cuerpo["format"] == "json" and cuerpo["stream"] and cuerpo["messages"][0]["role"] == "system"
        assert prov.modelos() == ["qwen3-vl:8b"]
        with pytest.raises(P.ErrorIA):
            P.crear(_cfg(srv, modelo="otro:1b")).chat([{"role": "user", "content": "x"}])
        url = srv.url
    with pytest.raises(P.ErrorIA):
        P.crear(P.Config.de({"proveedor": "ollama", "url": url})).modelos()
    assert P.crear({"proveedor": "plantilla"}) is None and P.crear({"proveedor": "anthropic"}) is None


# ── consulta ──
def test_consulta_cita_el_plano_marca_cifras_y_aprende(tmp_path):
    respuesta = "El cruce [C-1] es entre [AGUA-1] y [ALCANTARILLADO-2] con 0.5 ft; además 9.9 ft. [X-9]"
    with OllamaFalso(lambda c: respuesta) as srv:
        mem = M.Memoria(str(tmp_path / "P - IA"))
        prov = P.crear(_cfg(srv))
        r = consulta.preguntar(prov, "¿qué pasa en C-1?", _datos(), mem)
        assert r.referencias == ["C-1", "AGUA-1", "ALCANTARILLADO-2"] and r.sin_respaldo == [9.9]
        assert len(mem.registros(M.CONVERSACIONES)) == 1
        consulta.corregir(mem, r, "La diferencia es de 0.5 ft: hay que bajar el agua")
        pedido, usadas = consulta.armar_pedido("¿qué pasa en C-1?", _datos(), mem)
        assert usadas == 1 and "hay que bajar el agua" in pedido
        consulta.preguntar(prov, "¿y el drenaje?", _datos(), mem, historial=[(r.pregunta, r.texto)])
        msgs = srv.pedidos[-1]["messages"]
        assert [m["role"] for m in msgs] == ["system", "user", "assistant", "user"]


# ── informe redactado por la IA ──
def _respuesta_informe(cifra_falsa=None):
    def responder(c):
        msg = c["messages"][-1]["content"]
        if "resumen ejecutivo" in msg:
            return json.dumps({"resumen": "Resumen de la IA."})
        ids = [ln.split('"id": "')[1].split('"')[0] for ln in msg.splitlines() if '"id": "' in ln]
        fichas = {i: f"[{i}] redactada por la IA." for i in ids}
        if cifra_falsa and cifra_falsa in fichas:
            fichas[cifra_falsa] = "Separación de 99.9 ft."
        return json.dumps({"zona": "Zona revisada.", "fichas": fichas})
    return responder


def test_redactor_ia_revisa_cifras_recuerda_y_cae_a_la_plantilla(tmp_path):
    inf = _informe()
    with OllamaFalso(_respuesta_informe(cifra_falsa="N-1")) as srv:
        mem = M.Memoria(str(tmp_path / "P - IA"))
        red, motivo = redactor_ia.crear(_cfg(srv), mem)
        assert isinstance(red, redactor_ia.RedactorIA) and not motivo
        red.verificar()
        cuenta = red.preparar(inf)
        f = {x["id"]: x for x in inf.fichas}
        assert red.ficha(f["C-1"]) == "C-1 redactada por la IA."               # sin corchetes en el Word
        assert "Qué hacer" in red.ficha(f["N-1"])                                # 99.9 no sale de la ficha
        assert cuenta["ia"] == 1 and cuenta["revisadas"] == 1
        assert red.resumen(inf) == "Resumen de la IA." and red.zona(inf.zonas[0]) == "Zona revisada."
        n = len(srv.pedidos)
        red2, _ = redactor_ia.crear(_cfg(srv), M.Memoria(str(tmp_path / "P - IA")))
        red2.preparar(inf)
        assert red2.ficha(f["C-1"]) == "C-1 redactada por la IA." and red2.cuenta["memoria"] == 1
        assert len(srv.pedidos) - n == 1                                         # solo la ficha rechazada
        with pytest.raises(P.ErrorIA):
            redactor_ia.crear(_cfg(srv, modelo="otro:1b"), mem)[0].verificar()
    assert redactor_ia.crear({"proveedor": "plantilla"})[0].nombre == "plantilla"


# ── en la ventana ──
def _ventana(srv, tmp_path):
    from PySide6 import QtWidgets
    QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    from ui.ventana.app_window import Main
    proy = str(tmp_path / "Escenarios.digproj")
    shutil.copy(ESC, proy)
    w = Main()
    w._info = lambda *a, **k: None
    w._config_ia = lambda: {"proveedor": "ollama", "url": srv.url, "modelo": "qwen3-vl:8b", "paralelo": 2}
    w._open_project_path(proy)
    w._resultados_al_dia()
    return w, proy


def test_preguntale_al_proyecto_en_la_ventana(tmp_path):
    from PySide6 import QtCore, QtWidgets
    with OllamaFalso(lambda c: "Hay que revisar [C-1] y [C-2].") as srv:
        w, proy = _ventana(srv, tmp_path)
        dlg = w.abrir_pregunta()
        dlg.enviar("¿Qué conflictos hay?")
        fin = time.time() + 20
        while dlg._activo is not None and time.time() < fin:
            QtWidgets.QApplication.processEvents()
            time.sleep(0.01)
        tu = dlg.turnos[-1]
        assert tu["estado"] == "listo" and tu["res"].referencias == ["C-1", "C-2"]
        assert "href='ref:C-1'" in dlg.vista.toHtml().replace('"', "'") or "ref:C-1" in dlg.vista.toHtml()
        assert w._ia_ir_a("C-1")
        dlg._enlace(QtCore.QUrl("util:0"))
        carpeta = os.path.splitext(proy)[0] + " - IA"
        assert {"conversaciones.jsonl", "ensenanzas.jsonl", "LEEME.txt"} <= set(os.listdir(carpeta))
        w._dirty = False


def test_informe_con_ia_y_su_barra(tmp_path, monkeypatch):
    from ui.comun import barra_pesada
    vistos = []
    original = barra_pesada.AvancePesado._mostrar
    monkeypatch.setattr(barra_pesada.AvancePesado, "_mostrar",
                        lambda self, pump: vistos.append(self.barra.milesimas) or original(self, pump))
    import docx
    with OllamaFalso(_respuesta_informe()) as srv:
        w, _proy = _ventana(srv, tmp_path)
        ruta = str(tmp_path / "inf.docx")
        vistos.clear()
        assert w.generar_informe(ruta, Opciones(True, True, True, False)) == ruta
        txt = "\n".join(p.text for p in docx.Document(ruta).paragraphs)
        assert "Resumen de la IA." in txt and "redactada por la IA" in txt and "Redactado con IA" in txt
        assert vistos == sorted(vistos) and vistos[-1] == 1000
        n = len(srv.pedidos)
        w.generar_informe(ruta, Opciones(True, True, True, False))
        assert len(srv.pedidos) == n                                             # todo de la memoria
        url = srv.url
        w._dirty = False
    # servidor apagado: el informe sale igual, con plantillas y el motivo
    w.generar_informe(ruta, Opciones(True, True, True, False))
    assert url and "redactada por la IA" in "\n".join(p.text for p in docx.Document(ruta).paragraphs)
