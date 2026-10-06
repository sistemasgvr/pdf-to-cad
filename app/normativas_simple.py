"""Formato SIMPLE de normativas en Excel (el que propusieron los ingenieros civiles,
2026-10-06): una hoja «All (Flat)» con UNA FILA POR REGLA, en inglés, legible sin la
app — Utility A / Utility B / Orientation / Case / Min / Max (valor + unidad ft|in) /
Measured From / Notes / Reference(s) / Source Sheet — más «References» (texto de cada
referencia «HOJA:n»), una pestaña por utilidad (sus filas, para leer) y la leyenda.
PURO (solo openpyxl).

Es el formato OFICIAL para cada normativa nueva: `exportar` lo escribe y
`normativas_excel.importar` lo reconoce (hoja «All (Flat)»). La plantilla anterior y
las «clearance tables» se siguen pudiendo importar.

Al importar solo se leen «All (Flat)», «References» y, si están, las hojas propias de
la app («Fitting Angles», «App Notes»); las pestañas por utilidad, la leyenda y el
anexo de cortes son copias o dibujos para leer. Columnas opcionales a la derecha
(grises): Rule ID, Active, Mandatory, Status, Note IDs — las escribe la app para que
el ida y vuelta no pierda nada; un ingeniero puede dejarlas vacías.
"""
from __future__ import annotations

import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import normativas_clearance as CL

HOJA = "All (Flat)"
HOJA_REFS = "References"
HOJA_ANGULOS = "Fitting Angles"
HOJA_NOTAS = "App Notes"

UTILIDAD_EN = {"AGUA": "WATER", "ALCANTARILLADO": "SEWER", "DRENAJE": "STORMDRAIN",
               "ELECTRICO": "ELECTRICAL", "GAS": "GAS", "TELECOM": "TELECOM"}
CONTRA_EN = dict(UTILIDAD_EN, VIA_FERREA="RAILROAD TRACK", BORDILLO="CURB & GUTTER",
                 SUMIDERO="CATCH BASIN", BUZON="MAINTENANCE HOLE", SUPERFICIE="COVER",
                 OTRA_TUBERIA="OTHER PIPELINE")
_ALIAS = {"TELECOMMUNICATIONS": "TELECOM", "COMMUNICATIONS": "TELECOM", "STORM DRAIN": "DRENAJE",
          "ELECTRIC": "ELECTRICO", "SANITARY SEWER": "ALCANTARILLADO", "MANHOLE": "BUZON",
          "CURB AND GUTTER": "BORDILLO", "TRACK": "VIA_FERREA", "SURFACE": "SUPERFICIE"}
ACCESORIO_EN = {"codo": "Bend", "tee": "Tee", "wye": "Wye", "cruz": "Cross"}

# (clave, encabezado, ancho, opcional-de-la-app)
COLUMNAS = [
    ("a", "Utility A", 14, False), ("b", "Utility B", 18, False), ("orient", "Orientation", 12, False),
    ("caso", "Case / Sub-type", 34, False), ("min", "Min (value)", 10, False), ("min_u", "Min (unit)", 9, False),
    ("max", "Max (value)", 10, False), ("max_u", "Max (unit)", 9, False), ("medido", "Measured From", 34, False),
    ("notas", "Notes", 60, False), ("refs", "Reference(s)", 14, False), ("hoja", "Source Sheet", 11, False),
    ("id", "Rule ID", 12, True), ("activa", "Active", 8, True), ("obligatoria", "Mandatory", 10, True),
    ("estado", "Status", 9, True), ("nota_ids", "Note IDs", 12, True),
]
COLUMNAS_ANGULOS = [("util", "Utility", 16), ("acc", "Fitting", 10), ("angulos", "Allowed angles (deg)", 22),
                    ("tol", "Tolerance (deg)", 14), ("notas", "Notes", 50), ("id", "Rule ID", 12),
                    ("activa", "Active", 8), ("obligatoria", "Mandatory", 10)]

_AZUL = PatternFill("solid", fgColor="1F4E78")
_GRIS = PatternFill("solid", fgColor="7F7F7F")
_AMBAR = PatternFill("solid", fgColor="FFF2CC")
_CAB = Font(name="Arial", color="FFFFFF", bold=True)
_TXT = Font(name="Arial")


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").upper().replace("_", " ")).strip()


def _codigo(nombre):
    """Texto de Utility A/B → código de la app (o None si no se conoce)."""
    n = _norm(nombre)
    for cod, en in CONTRA_EN.items():
        if n == en or n == cod.replace("_", " "):
            return cod
    return _ALIAS.get(n)


def es_formato_simple(wb):
    return HOJA in wb.sheetnames


# ─────────────────────────────── exportar ───────────────────────────────


def _en_unidad(valores, unidad=None):
    """Pies → (valores, unidad). Respeta la unidad original; si no, pulgadas cuando
    algún valor no es un pie entero y todos son pulgadas enteras (18" en vez de 1.5')."""
    vals = [v for v in valores if v is not None]
    if unidad not in ("ft", "in"):
        unidad = "in" if vals and any(abs(v - round(v)) > 1e-9 for v in vals) \
            and all(abs(v * 12 - round(v * 12)) < 1e-6 for v in vals) else "ft"
    f = 12.0 if unidad == "in" else 1.0
    out = [None if v is None else round(v * f, 4) for v in valores]
    return [int(v) if v is not None and v == int(v) else v for v in out], (unidad if vals else None)


def _fila(regla):
    p = regla.get("params") or {}
    tipo = regla.get("tipo")
    (mn, mx), u = _en_unidad([p.get("minimo"), p.get("maximo")], p.get("unidad"))
    contra = p.get("contra", "")
    caso = " -- ".join(x for x in (p.get("condicion"), p.get("condicion_contra")) if x) or "General"
    notas = p.get("descripcion") or regla.get("descripcion") or ""
    if tipo == "requisito" and not notas:
        notas = p.get("texto_original", "")
    return {
        "a": "; ".join(UTILIDAD_EN.get(x, x) for x in regla.get("utilidades") or []),
        "b": CONTRA_EN.get(contra, contra),
        "orient": "Horizontal" if tipo == "separacion_horizontal" else "Vertical"
        if tipo in ("separacion_vertical", "recubrimiento") else "",
        "caso": caso, "min": mn, "min_u": u if mn is not None else None,
        "max": mx, "max_u": u if mx is not None else None,
        "medido": p.get("medido_desde", ""), "notas": notas,
        "refs": ", ".join(p.get("referencias") or []), "hoja": p.get("hoja", ""),
        "id": regla["id"], "activa": "Yes" if regla.get("activa", True) else "No",
        "obligatoria": "Yes" if regla.get("obligatoria", True) else "No",
        "estado": "Review" if p.get("estado") == "Revisar" else "OK",
        "nota_ids": ", ".join(p.get("notas") or []),
    }


def _cabecera(ws, cols, fila=1):
    for c, (_k, h, w, *op) in enumerate(cols, 1):
        cel = ws.cell(row=fila, column=c, value=h)
        cel.font = _CAB
        cel.fill = _GRIS if op and op[0] else _AZUL
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=fila + 1, column=1).coordinate


def _volcar(ws, filas, desde):
    for f, datos in enumerate(filas, desde):
        for c, (k, *_r) in enumerate(COLUMNAS, 1):
            cel = ws.cell(row=f, column=c, value=datos.get(k))
            cel.font = _TXT
            cel.alignment = Alignment(wrap_text=k in ("notas", "caso", "medido"), vertical="top")
            if k == "estado" and datos.get(k) == "Review":
                cel.fill = _AMBAR


def exportar(reglas, anexos, ruta, version_app=""):
    """Escribe el libro en el formato simple con TODAS las reglas actuales."""
    anexos = anexos or {}
    distancias = [r for r in reglas if r.get("tipo") != "angulos_accesorio"]
    angulos = [r for r in reglas if r.get("tipo") == "angulos_accesorio"]
    filas = sorted((_fila(r) for r in distancias),
                   key=lambda d: (d["a"], d["orient"], d["b"], d["caso"]))
    wb = openpyxl.Workbook()
    _leyenda(wb.active, version_app)
    _como_agregar(wb.create_sheet("How to Add a Utility"))
    ws = wb.create_sheet(HOJA)
    _cabecera(ws, COLUMNAS)
    _volcar(ws, filas, 2)
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(COLUMNAS))}{len(filas) + 1}"
    for util in dict.fromkeys(d["a"] for d in filas):
        if not util or len(util) > 31 or any(ch in util for ch in "[]:*?/\\"):
            continue
        wu = wb.create_sheet(util)
        wu.cell(row=1, column=1, value=f"{util} -- copy of its rows in \"{HOJA}\" (for reading; "
                                         f"edit \"{HOJA}\": it is the sheet the app imports).").font = Font(name="Arial", bold=True)
        _cabecera(wu, COLUMNAS, fila=2)
        _volcar(wu, [d for d in filas if d["a"] == util], 3)

    wr = wb.create_sheet(HOJA_REFS)
    _cabecera(wr, [("s", "Source Sheet", 12), ("n", "Ref #", 8), ("t", "Standard / Note Text", 100)])
    for f, (rid, ref) in enumerate(sorted((anexos.get("referencias") or {}).items()), 2):
        texto = ref.get("texto", "") if isinstance(ref, dict) else str(ref)
        hoja, num = rid.rsplit(":", 1) if ":" in rid else ("", rid)
        for c, v in enumerate((hoja, num, texto), 1):
            wr.cell(row=f, column=c, value=v).alignment = Alignment(wrap_text=c == 3, vertical="top")

    if angulos:
        wa = wb.create_sheet(HOJA_ANGULOS)
        _cabecera(wa, COLUMNAS_ANGULOS)
        for f, r in enumerate(angulos, 2):
            p = r.get("params") or {}
            vals = ("; ".join(UTILIDAD_EN.get(x, x) for x in r.get("utilidades") or []),
                    ACCESORIO_EN.get(p.get("accesorio"), p.get("accesorio", "")),
                    "; ".join(f"{a:g}" for a in p.get("angulos") or []), p.get("tolerancia"),
                    p.get("descripcion") or r.get("descripcion") or "", r["id"],
                    "Yes" if r.get("activa", True) else "No", "Yes" if r.get("obligatoria", True) else "No")
            for c, v in enumerate(vals, 1):
                wa.cell(row=f, column=c, value=v)
    if anexos.get("notas"):
        wn = wb.create_sheet(HOJA_NOTAS)
        _cabecera(wn, [("i", "ID", 10), ("s", "Source Sheet", 14), ("a", "Applies to", 16), ("t", "Text", 100)])
        for f, n in enumerate(anexos["notas"], 2):
            for c, v in enumerate((n.get("id"), n.get("hoja"), n.get("aplica"), n.get("texto")), 1):
                wn.cell(row=f, column=c, value=v).alignment = Alignment(wrap_text=c == 4, vertical="top")
    wb.active = 2
    wb.save(ruta)


def _texto(ws, lineas, ancho_b=110):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = ancho_b
    for f, (a, b) in enumerate(lineas, 1):
        ws.cell(row=f, column=1, value=a).font = Font(name="Arial", bold=True, size=13 if f == 1 else 10)
        if b:
            ws.cell(row=f, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")


def _leyenda(ws, version_app):
    ws.title = "Legend"
    _texto(ws, [
        ("Purpose", "Design clearances, one row per utility pair, so any rule can be found and filtered "
                    f"without the app. Exported by PDF-to-CAD {version_app}.".strip()),
        ("Utility A / Utility B", "Utility A owns the rule; Utility B is the other utility or element (WATER, SEWER, "
                                  "STORMDRAIN, ELECTRICAL, GAS, TELECOM, RAILROAD TRACK, CURB & GUTTER, COVER, "
                                  "CATCH BASIN, MAINTENANCE HOLE). Both directions are kept separately."),
        ("Orientation", "Horizontal or Vertical. Utility B = COVER means depth of cover."),
        ("Case / Sub-type", "Condition the value applies to (diameter, pressure, main vs. service…). \"General\" = "
                            "a single value."),
        ("Min / Max", "Number + unit (ft or in). Max only when the source gives a range. A row with no number "
                      "is a written requirement: put the text in Notes."),
        ("Measured From", "Physical reference point, when the source names it."),
        ("Reference(s)", "\"SHEET:n\" codes, separated by commas; their text goes on the \"References\" tab."),
        ("Grey columns", "Optional, filled by the app (Rule ID, Active, Mandatory, Status, Note IDs). "
                         "Leave them empty in new rows."),
        ("Import", f"The app reads \"{HOJA}\" and \"References\" (plus \"{HOJA_ANGULOS}\"/\"{HOJA_NOTAS}\" if "
                   "present). Per-utility tabs are copies for reading."),
    ])


def _como_agregar(ws):
    _texto(ws, [
        ("How to add rules", None),
        ("1. Name", "Use the utility labels exactly as listed in the Legend (UPPERCASE)."),
        ("2. Add rows", f"Add one row per rule to \"{HOJA}\": Utility A, Utility B, Orientation, Case "
                        "(or \"General\"), Min value + unit, Max only for a range, Measured From, Notes, "
                        "Reference(s), Source Sheet."),
        ("3. Other direction", "If the other utility's standard gives its own value, add it as a separate row "
                               "with A and B swapped. Do not assume symmetry."),
        ("4. References", "Pick a short code for the source document (e.g. RW-1) and cite \"RW-1:3\"; add a "
                          "row to \"References\" with Source Sheet, Ref # and the full text."),
        ("5. Import", "In the app: Design standards → Import Excel."),
    ])


# ─────────────────────────────── importar ───────────────────────────────


def _num(v):
    if v in (None, ""):
        return None
    return float(str(v).replace(",", ".").strip())


def _a_pies(valor, unidad, f):
    """(pies, unidad, error)."""
    try:
        n = _num(valor)
    except ValueError:
        return None, None, f"«{valor}» is not a number."
    if n is None:
        return None, None, None
    u = _norm(unidad).lower().strip(".'\"") or "ft"
    u = {"feet": "ft", "foot": "ft", "inch": "in", "inches": "in", "pies": "ft", "pulg": "in"}.get(u, u)
    if u not in ("ft", "in"):
        return None, None, f"Unidad no válida: «{unidad}» (ft o in)."
    if n < 0:
        return None, None, "Las distancias no pueden ser negativas."
    return (n / 12.0 if u == "in" else n), u, None


def _si(v):
    return v in (None, "") or _norm(v) in ("YES", "Y", "SI", "SÍ", "TRUE", "1", "X")


def _refs(wb):
    refs = {}
    if HOJA_REFS in wb.sheetnames:
        for fila in wb[HOJA_REFS].iter_rows(min_row=2, values_only=True):
            if not fila or len(fila) < 3 or fila[1] in (None, ""):
                continue
            hoja, num = str(fila[0] or "").strip(), str(fila[1]).strip()
            num = num[:-2] if num.endswith(".0") else num
            refs[f"{hoja}:{num}" if hoja else num] = {"texto": str(fila[2] or ""), "origen": hoja}
    return refs


def _notas(wb):
    if HOJA_NOTAS not in wb.sheetnames:
        return []
    return [{"id": str(f[0]).strip(), "hoja": str(f[1] or ""), "aplica": str(f[2] or ""), "texto": str(f[3] or "")}
            for f in wb[HOJA_NOTAS].iter_rows(min_row=2, values_only=True) if f and f[0]]


def _leer(ws, cols):
    cab = [_norm(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    pos = {k: cab.index(_norm(h)) for k, h, *_r in cols if _norm(h) in cab}
    for f, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        v = {k: (fila[i] if i < len(fila) else None) for k, i in pos.items()}
        if any(x not in (None, "") for x in v.values()):
            yield f, v, pos


def importar(wb, archivo=""):
    """(reglas, anexos, errores) del libro en formato simple."""
    refs = _refs(wb)
    anexos = {"referencias": refs, "notas": _notas(wb),
              "documento": {"archivo": archivo, "titulos": [], "tablas": []}}
    reglas, errores, ids, cuenta = [], [], set(), {}
    ws = wb[HOJA]
    for f, v, pos in _leer(ws, COLUMNAS):
        if "a" not in pos or "b" not in pos:
            return [], anexos, [(1, f"Faltan las columnas «Utility A» y «Utility B» en «{HOJA}».")]
        regla, err = _regla(v, refs, cuenta)
        if not err and regla["id"] in ids:
            err = f"El ID «{regla['id']}» está repetido."
        if err:
            errores.append((f, err))
            continue
        ids.add(regla["id"])
        reglas.append(regla)
    if HOJA_ANGULOS in wb.sheetnames:
        for f, v, _pos in _leer(wb[HOJA_ANGULOS], COLUMNAS_ANGULOS):
            regla, err = _regla_angulo(v, f)
            if not err and regla["id"] in ids:
                err = f"El ID «{regla['id']}» está repetido."
            if err:
                errores.append((f, f"{HOJA_ANGULOS}: {err}"))
                continue
            ids.add(regla["id"])
            reglas.append(regla)
    return reglas, anexos, errores


def _regla(v, refs, cuenta):
    utils = []
    for u in re.split(r"[;,]", str(v.get("a") or "")):
        if not u.strip():
            continue
        cod = _codigo(u)
        if cod not in UTILIDAD_EN:
            return None, f"Utility A no válida: «{u.strip()}»."
        utils.append(cod)
    if not utils:
        return None, "Falta Utility A."
    b_txt = str(v.get("b") or "").strip()
    if not b_txt:
        return None, "Falta Utility B."
    contra = _codigo(b_txt)
    estado = "OK" if contra else "Revisar"                  # elemento que la app no conoce: se guarda tal cual
    contra = contra or b_txt
    mn, u1, e1 = _a_pies(v.get("min"), v.get("min_u"), 0)
    mx, u2, e2 = _a_pies(v.get("max"), v.get("max_u") or v.get("min_u"), 0)
    if e1 or e2:
        return None, e1 or e2
    if mn is not None and mx is not None and mx < mn:
        return None, "El máximo es menor que el mínimo."
    orient = _norm(v.get("orient"))
    if mn is None and mx is None:
        tipo = "requisito"
    elif contra == "SUPERFICIE":
        tipo = "recubrimiento"
    elif orient.startswith("H"):
        tipo = "separacion_horizontal"
    elif orient.startswith("V"):
        tipo = "separacion_vertical"
    else:
        return None, f"Orientation no válida: «{v.get('orient') or ''}» (Horizontal o Vertical)."
    citadas = [r.strip() for r in re.split(r"[,;]", str(v.get("refs") or "")) if r.strip()]
    if any(r not in refs for r in citadas):
        estado = "Revisar"
    if _norm(v.get("estado")) in ("REVIEW", "REVISAR"):
        estado = "Revisar"
    hoja = str(v.get("hoja") or "").strip()
    rid = str(v.get("id") or "").strip()
    if not rid:
        letra = {"separacion_horizontal": "H", "separacion_vertical": "V", "recubrimiento": "C", "requisito": "R"}[tipo]
        pref = hoja or CL.ABREV.get(utils[0], utils[0][:2])
        cuenta[(pref, letra)] = cuenta.get((pref, letra), 0) + 1
        rid = f"{pref}-{letra}-{cuenta[(pref, letra)]:02d}"
    caso = str(v.get("caso") or "").strip()
    notas_txt = str(v.get("notas") or "").strip()
    return {"id": rid, "tipo": tipo, "utilidades": utils, "activa": _si(v.get("activa")),
            "obligatoria": _si(v.get("obligatoria")),
            "params": {
                "contra": contra, "condicion": "" if _norm(caso) == "GENERAL" else caso, "condicion_contra": "",
                "minimo": mn, "maximo": mx, "unidad": u1 or u2,
                "medido_desde": str(v.get("medido") or "").strip(),
                "referencias": citadas,
                "notas": [n.strip() for n in re.split(r"[,;]", str(v.get("nota_ids") or "")) if n.strip()],
                "descripcion": notas_txt, "estado": estado,
                "texto_original": notas_txt if tipo == "requisito" else "",
                "encabezado_original": b_txt, "condicion_original": caso, "hoja": hoja,
            }}, None


def _regla_angulo(v, f):
    utils = [_codigo(u) for u in re.split(r"[;,]", str(v.get("util") or "")) if u.strip()]
    if not utils or any(u not in UTILIDAD_EN for u in utils):
        return None, f"Utility no válida: «{v.get('util') or ''}»."
    acc = next((k for k, en in ACCESORIO_EN.items() if _norm(en) == _norm(v.get("acc"))), None)
    if not acc:
        return None, "Fitting debe ser Bend, Tee, Wye o Cross."
    try:
        angulos = sorted({round(float(a.replace(",", ".")), 4)
                          for a in re.split(r"[;\n]", str(v.get("angulos") or "")) if a.strip()})
        tol = _num(v.get("tol"))
    except ValueError:
        return None, "Ángulos o tolerancia no es un número."
    if not angulos or any(not (0 < a <= 180) for a in angulos):
        return None, "Ángulos permitidos: al menos uno, entre 0 y 180."
    rid = str(v.get("id") or "").strip() or f"ANG-{f:03d}"
    return {"id": rid, "tipo": "angulos_accesorio", "utilidades": utils, "activa": _si(v.get("activa")),
            "obligatoria": _si(v.get("obligatoria")),
            "params": {"accesorio": acc, "angulos": angulos, "tolerancia": 1.0 if tol is None else tol,
                       "descripcion": str(v.get("notas") or "")}}, None
