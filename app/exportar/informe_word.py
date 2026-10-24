"""El INFORME AUTOMÁTICO en Word (sin Qt; usa python-docx). Pedido del usuario 2026-10-23.

`escribir(informe, redactor, ruta, imagenes)`: portada, resumen (texto del redactor + conteos), mapa general
con las zonas numeradas, una ficha por zona (captura 2D y/o 3D + cada conflicto y aviso explicado) y un anexo
con todo en una tabla. `imagenes` = {"mapa": bytes, ("2d", n): bytes, ("3d", n): bytes} (JPG o PNG).

Rápido con proyectos enormes: el anexo llega hasta `ANEXO_MAX` filas (una tabla de Word es lo más lento de
armar; el detalle completo está en el Excel de la Tabla de datos) y avisa su avance por `nucleo.avance`
(etapas «informe_zonas», «informe_anexo», «informe_guardar»), como las demás barras precisas.
"""
from __future__ import annotations

import io

from nucleo import avance
from traduccion.i18n_core import t

ANEXO_MAX = 1000
ANCHO_IMAGEN_CM = 16.0
COLOR = {"ERROR": (192, 0, 0), "AVISO": (191, 128, 0), "INFO": (31, 111, 179)}
NIVEL = {"ERROR": "Error", "AVISO": "Aviso", "INFO": "Información"}


def _rgb(c):
    from docx.shared import RGBColor
    return RGBColor(*c)


def _tabla(doc, cabecera, filas):
    from docx.shared import Pt
    tb = doc.add_table(rows=1, cols=len(cabecera))
    tb.style = "Light Grid Accent 1"
    for c, txt in zip(tb.rows[0].cells, cabecera):
        c.text = txt
    for fila in filas:
        celdas = tb.add_row().cells
        for c, txt in zip(celdas, fila):
            c.text = "" if txt is None else str(txt)
    for fila in tb.rows:
        for c in fila.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return tb


def _imagen(doc, datos, pie):
    from docx.shared import Cm, Pt
    if not datos:
        return
    doc.add_picture(io.BytesIO(datos), width=Cm(ANCHO_IMAGEN_CM))
    p = doc.add_paragraph()
    r = p.add_run(pie)
    r.italic, r.font.size = True, Pt(9)


def _ficha(doc, f, redactor):
    from docx.shared import Pt
    p = doc.add_paragraph()
    r = p.add_run(f"{f['id']} · {f['titulo']}")
    r.bold = True
    estado = (t("aprobado") if f.get("aprobado") else t(NIVEL.get(f["nivel"], f["nivel"])).lower())
    r2 = p.add_run(f"  [{estado}]")
    r2.font.size = Pt(9)
    r2.font.color.rgb = _rgb((90, 140, 90) if f.get("aprobado") else COLOR.get(f["nivel"], (80, 80, 80)))
    doc.add_paragraph(redactor.ficha(f))


def _estado(f):
    if f.get("aprobado"):
        return t("Aprobado")
    return t("Información") if f["info"] else t("Pendiente")


def escribir(informe, redactor, ruta, imagenes=None, notas=()):
    """Escribe el .docx. `notas` = frases que van en el resumen (p. ej. «no hubo capturas 3D: …»)."""
    import docx
    from docx.shared import Pt
    imagenes = imagenes or {}
    doc = docx.Document()
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Arial", Pt(10)

    doc.add_heading(t("Informe de conflictos y normativas"), 0)
    if informe.titulo:
        doc.add_paragraph(t("Proyecto: {p}").format(p=informe.titulo))
    doc.add_paragraph(t("Fecha: {f}").format(f=informe.fecha))
    contenido = [x for x, on in ((t("conflictos"), informe.opciones.conflictos),
                                 (t("normativas"), informe.opciones.normativas),
                                 (t("avisos de reconocimiento"), bool(informe.resumen.get("reconocimiento")))) if on]
    doc.add_paragraph(t("Contenido: {c}.").format(c=(" " + t("y") + " ").join(contenido)))

    doc.add_heading(t("Resumen"), 1)
    doc.add_paragraph(redactor.resumen(informe))
    for n in notas:
        doc.add_paragraph(n).runs[0].italic = True
    r = informe.resumen
    filas = [(t("Utilidades revisadas"), r["utilidades"])]
    if informe.opciones.conflictos:
        filas.append((t("Conflictos"), r["conflictos"]))
    if informe.opciones.normativas:
        filas += [(t("Avisos de normativa pendientes"), r["avisos_pendientes"]), (t("Avisos aprobados"), r["aprobados"]),
                  (t("Avisos informativos"), r["info"]), (t("Utilidades sin tipo"), r["sin_tipo"])]
    if r.get("reconocimiento"):
        filas.append((t("Avisos de reconocimiento (diámetro del plano)"), r["reconocimiento"]))
    _tabla(doc, [t("Concepto"), t("Cantidad")], filas)
    if r["por_titulo"]:
        doc.add_paragraph()
        _tabla(doc, [t("Qué"), t("Cantidad")], [(tit, n) for (_tipo, tit), n in r["por_titulo"]])
    if r["por_utilidad"]:
        doc.add_paragraph()
        from nucleo.informe_redaccion import nombre_utilidad
        _tabla(doc, [t("Utilidad"), t("Conflictos y avisos")], [(nombre_utilidad(c), n) for c, n in r["por_utilidad"]])

    if imagenes.get("mapa"):
        doc.add_heading(t("Mapa general"), 1)
        _imagen(doc, imagenes["mapa"], t("Todo el dibujo con las zonas numeradas."))

    if informe.zonas:
        doc.add_heading(t("Fichas por zona"), 1)
    n = len(informe.zonas)
    for k, z in enumerate(informe.zonas, 1):
        avance.avisar("informe_zonas", k, n)
        doc.add_heading(t("Zona {z}").format(z=z.numero), 2)
        doc.add_paragraph(redactor.zona(z))
        _imagen(doc, imagenes.get(("2d", z.numero)), t("Zona {z} en el plano.").format(z=z.numero))
        _imagen(doc, imagenes.get(("3d", z.numero)), t("Zona {z} en la vista 3D.").format(z=z.numero))
        for f in z.fichas:
            _ficha(doc, f, redactor)

    rec = [f for f in informe.fichas if f["tipo"] == "reconocimiento"]
    if rec:
        doc.add_heading(t("Avisos de reconocimiento"), 1)
        doc.add_paragraph(t("Lo que el reconocimiento del PDF no pudo leer con certeza (por ejemplo, el diámetro "
                            "que dice el plano). No son incumplimientos de normativa: revísalos en el plano."))
        _tabla(doc, [t("ID"), t("Utilidades"), t("Qué")],
               [(f["id"], ", ".join(u["codigo"] for u in f["utilidades"]), f["mensaje_app"] or f["titulo"])
                for f in rec[:ANEXO_MAX]])

    if informe.fichas:
        doc.add_heading(t("Anexo: todos los conflictos y avisos"), 1)
        filas = []
        total = min(len(informe.fichas), ANEXO_MAX)
        for i, f in enumerate(informe.fichas[:ANEXO_MAX], 1):
            if i % 100 == 0 or i == total:
                avance.avisar("informe_anexo", i, total)
            filas.append((f["id"], f["titulo"], t(NIVEL.get(f["nivel"], f["nivel"])),
                          ", ".join(u["codigo"] for u in f["utilidades"]), f.get("zona", ""), _estado(f)))
        _tabla(doc, [t("ID"), t("Qué"), t("Nivel"), t("Utilidades"), t("Zona"), t("Estado")], filas)
        if len(informe.fichas) > ANEXO_MAX:
            doc.add_paragraph(t("Se muestran {n} de {total}. El detalle completo está en la Tabla de datos "
                                "(Herramientas → Tabla de datos → Exportar a Excel).").format(
                n=ANEXO_MAX, total=len(informe.fichas)))
    avance.avisar("informe_guardar")
    doc.save(ruta)
    avance.avisar("informe_guardar", 1, 1)
    return ruta
