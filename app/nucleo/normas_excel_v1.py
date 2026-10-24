"""Lectura del Excel de normativas VIEJO, de cuatro hojas (PURO; usa openpyxl). Movido tal cual de
`normas_excel` el 2026-10-21, al pasar a la tabla única de reglas: un Excel viejo (TIPOS, PRESION-DIAMETROS,
PRESION-ACCESORIOS, ELECT-TELECOM) se sigue importando —`normas_excel.importar` lo detecta— y `normas_migrar`
lo pasa a reglas. Los encabezados se buscan por NOMBRE; `importar(ruta, problemas=[])` anota cada celda que no
se entiende (hoja, fila, columna) y no la usa; las hojas que no son de las cuatro tablas nunca se leen.
"""
from __future__ import annotations

from nucleo import normas_catalogo as nc

HOJA_TIPOS = "TIPOS"
HOJA_DIAM = "PRESION-DIAMETROS"
HOJA_ACC = "PRESION-ACCESORIOS"
HOJA_ELEC = "ELECT-TELECOM"

COL_TIPOS = ("UTILIDAD", "TIPO")
COL_DIAM = ("UTILIDAD", "TIPO", "DIAMETRO (in)", "DIAMETRO MIN (in)", "DIAMETRO MAX (in)", "SOLO SE PERMITE EN")
COL_ACC = ("UTILIDAD", "ACCESORIO", "SOLO SE PERMITE EN", "DIAMETRO PRINCIPAL (in)", "PROHIBIDO USO",
           "DIAMETRO RAMAL DISTRIBUCION", "ANGULOS PERMITIDOS")
COL_ELEC = ("UTILIDAD", "ESTRUCTURA", "TIPO", "AMPERAJE MIN (A)", "AMPERAJE MAX (A)", "LONGITUD MIN (pies)",
            "LONGITUD MAX (pies)", "DIAMETRO FIJO (in)", "DIAMETRO MIN (in)", "DIAMETRO MAX (in)",
            "RADIO CURVA MIN (pies)", "RADIO CURVA MAX (pies)", "RADIO CURVA FIJO (pies)")
# Columnas de UN número (validación numérica al exportar; aviso al importar si no lo es).
NUMERICAS = {"DIAMETRO MIN (in)", "DIAMETRO MAX (in)", "AMPERAJE MIN (A)", "AMPERAJE MAX (A)",
             "LONGITUD MIN (pies)", "LONGITUD MAX (pies)", "RADIO CURVA MIN (pies)", "RADIO CURVA MAX (pies)",
             "DIAMETRO RAMAL DISTRIBUCION"}
LISTAS = {"DIAMETRO (in)", "DIAMETRO PRINCIPAL (in)", "DIAMETRO FIJO (in)", "ANGULOS PERMITIDOS",
          "RADIO CURVA FIJO (pies)"}
# Otros nombres que se aceptan para una columna (el ingeniero le sumó «(in)», 2026-10-14).
ALIAS = {"DIAMETRO PRINCIPAL (in)": ("DIAMETRO PRINCIPAL",)}
# Columnas del formato viejo que se siguen leyendo (rango «0-320» en una celda).
VIEJAS = {HOJA_ELEC: ("AMPERAJE (A)", "LONGITUD RANGO (pies)")}


# ─────────────────────────── importar ───────────────────────────

def _hoja(wb, nombre):
    for ws in wb.worksheets:
        if nc.clave(ws.title) == nc.clave(nombre):
            return ws
    return None


def _filas(ws):
    """[{encabezado: valor}] de la primera tabla de la hoja (fila 1 = encabezados)."""
    if ws is None:
        return []
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    cab = [nc.clave(c) for c in filas[0]]
    out = []
    for n, f in enumerate(filas[1:], 2):
        d = {cab[i]: v for i, v in enumerate(f) if i < len(cab) and cab[i]}
        if any(v not in (None, "") for v in d.values()):
            d["_fila"] = n
            out.append(d)
    return out


class _Lector:
    """Lee celdas y anota las que no se entienden («Hoja, fila N, COLUMNA: «>=400» no es un número»)."""

    def __init__(self, problemas):
        self.problemas = problemas if problemas is not None else []

    def _mal(self, hoja, d, col, v, que):
        self.problemas.append(f"{hoja}, fila {d.get('_fila', '?')}, {col}: «{v}» {que}")

    def numero(self, hoja, d, col):
        v = _val(d, col)
        n = nc.numero(v)
        if n is None and v not in (None, "") and str(v).strip():
            self._mal(hoja, d, col, v, "no es un número")
        return n

    def lista(self, hoja, d, col):
        v = _val(d, col)
        nums = nc.lista_numeros(v)
        partes = [x for x in str(v).split(";") if x.strip()] if v not in (None, "") else []
        if len(nums) != len(partes) and not isinstance(v, (int, float)):
            self._mal(hoja, d, col, v, "no es una lista como 6; 8; 12")
        return nums

    def rango(self, hoja, d, col_min, col_max, col_vieja):
        """[mín, máx] de las columnas MIN/MAX, o del formato viejo «0-320» en una celda."""
        r = nc.rango_min_max(self.numero(hoja, d, col_min), self.numero(hoja, d, col_max))
        if r is None and col_vieja:
            v = d.get(nc.clave(col_vieja))
            r = nc.rango(v)
            if r is None and v not in (None, "") and str(v).strip():
                self._mal(hoja, d, col_vieja, v, "no se entiende (usa MIN y MAX)")
        return r


def _val(d, col):
    """Valor de la columna `col` (o de uno de sus otros nombres, `ALIAS`)."""
    for nombre in (col,) + ALIAS.get(col, ()):
        if nc.clave(nombre) in d:
            return d[nc.clave(nombre)]
    return None


def _txt(d, col):
    v = _val(d, col)
    return "" if v is None else str(v).strip()


def _columnas_desconocidas(ws, columnas, lee):
    """Anota cada encabezado de la tabla que la app no conoce (sus datos no se usarían)."""
    if ws is None:
        return
    conocidas = {nc.clave(c) for c in columnas}
    conocidas |= {nc.clave(a) for c in columnas for a in ALIAS.get(c, ())}
    conocidas |= {nc.clave(c) for c in VIEJAS.get(ws.title.upper(), ())}
    if nc.clave(ws.title) == nc.clave(HOJA_TIPOS):
        conocidas |= {nc.clave(u) for u in nc.UTILIDADES}          # bloque auxiliar oculto (K:P)
    fila = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    for v in fila:
        if v not in (None, "") and nc.clave(v) not in conocidas and not str(v).strip().startswith("_"):
            lee.problemas.append(f"{ws.title}: columna «{v}» desconocida")


def importar(ruta, problemas=None):
    """(catálogo, hojas que faltan). Si se pasa la lista `problemas`, se le agrega cada celda
    que no se entendió (y que por eso no se usa)."""
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True)
    lee = _Lector(problemas)
    hojas = {n: _hoja(wb, n) for n in (HOJA_TIPOS, HOJA_DIAM, HOJA_ACC, HOJA_ELEC)}
    avisos = [n for n, ws in hojas.items() if ws is None]
    for n, cols in ((HOJA_TIPOS, COL_TIPOS), (HOJA_DIAM, COL_DIAM), (HOJA_ACC, COL_ACC), (HOJA_ELEC, COL_ELEC)):
        _columnas_desconocidas(hojas[n], cols, lee)
    return armar({n: _filas(ws) for n, ws in hojas.items()}, lee.problemas), avisos


def armar(hojas, problemas=None, tipos_estrictos=False):
    """Catálogo de las filas de cada tabla: `hojas` = {HOJA: [{columna (clave): valor, "_fila": n}]}.
    Lo usan el Excel y el editor de la ventana (`normas_tablas`): las MISMAS reglas para leer cada celda.
    La aplicación manda: una fila cuya UTILIDAD no es de la app, o cuyo ACCESORIO no reconoce, no se usa
    (se anota en `problemas`); con `tipos_estrictos` (el editor), tampoco un TIPO que no está en TIPOS."""
    lee = _Lector(problemas)
    cat = {"tipos": [], "diametros": [], "accesorios": [], "electricas": []}

    def utilidad(hoja, d):
        u = nc.clave(_txt(d, "UTILIDAD"))
        if u in nc.UTILIDADES:
            return u
        lee._mal(hoja, d, "UTILIDAD", _txt(d, "UTILIDAD"), "no es una utilidad de la aplicación")
        return None

    for d in hojas.get(HOJA_TIPOS) or []:
        if not _txt(d, "UTILIDAD") and not _txt(d, "TIPO"):
            continue                       # solo el bloque auxiliar (K:P) tiene algo en esa fila
        u, t = utilidad(HOJA_TIPOS, d), nc.clave(_txt(d, "TIPO"))
        if u and t:
            cat["tipos"].append([u, t])
        elif u:
            lee._mal(HOJA_TIPOS, d, "TIPO", "", "está vacío")
    tipos = {tuple(x) for x in cat["tipos"]}

    def tipo(hoja, d, u, col):
        t = nc.clave(_txt(d, col))
        if t and tipos_estrictos and (u, t) not in tipos:
            lee._mal(hoja, d, col, _txt(d, col), "no es un tipo de {u} (hoja TIPOS)".format(u=u))
            return None
        return t

    H = HOJA_DIAM
    for d in hojas.get(H) or []:
        u = utilidad(H, d)
        t = tipo(H, d, u, "TIPO") if u else None
        if u and t is not None:
            cat["diametros"].append({"utilidad": u, "tipo": t,
                                     "diametros": lee.lista(H, d, "DIAMETRO (in)"),
                                     "diam_min": lee.numero(H, d, "DIAMETRO MIN (in)"),
                                     "diam_max": lee.numero(H, d, "DIAMETRO MAX (in)"),
                                     "nota": _txt(d, "SOLO SE PERMITE EN")})
    H = HOJA_ACC
    for d in hojas.get(H) or []:
        u = utilidad(H, d)
        if not u:
            continue
        if nc.ACCESORIOS.get(nc.clave(_txt(d, "ACCESORIO"))) is None:
            lee._mal(H, d, "ACCESORIO", _txt(d, "ACCESORIO"), "no es un accesorio que la aplicación reconoce")
            continue
        solo, prohibido = tipo(H, d, u, "SOLO SE PERMITE EN"), tipo(H, d, u, "PROHIBIDO USO")
        if solo is None or prohibido is None:
            continue
        cat["accesorios"].append({
            "utilidad": u, "accesorio": _txt(d, "ACCESORIO"), "solo_en": solo,
            "diam_principal": lee.lista(H, d, "DIAMETRO PRINCIPAL (in)"), "prohibido_en": prohibido,
            "ramal": lee.numero(H, d, "DIAMETRO RAMAL DISTRIBUCION"),
            "angulos": lee.lista(H, d, "ANGULOS PERMITIDOS")})
    H = HOJA_ELEC
    for d in hojas.get(H) or []:
        u = utilidad(H, d)
        t = tipo(H, d, u, "TIPO") if u else None
        if not u or t is None:
            continue
        cat["electricas"].append({
            "utilidad": u, "estructura": nc.clave(_txt(d, "ESTRUCTURA")), "tipo": t,
            "amperaje": lee.rango(H, d, "AMPERAJE MIN (A)", "AMPERAJE MAX (A)", "AMPERAJE (A)"),
            "longitud": lee.rango(H, d, "LONGITUD MIN (pies)", "LONGITUD MAX (pies)", "LONGITUD RANGO (pies)"),
            "diam_fijo": lee.lista(H, d, "DIAMETRO FIJO (in)"),
            "diam_min": lee.numero(H, d, "DIAMETRO MIN (in)"),
            "diam_max": lee.numero(H, d, "DIAMETRO MAX (in)"),
            "radio_min": lee.numero(H, d, "RADIO CURVA MIN (pies)"),
            "radio_max": lee.numero(H, d, "RADIO CURVA MAX (pies)"),
            "radio_fijo": lee.lista(H, d, "RADIO CURVA FIJO (pies)")})
    return cat
