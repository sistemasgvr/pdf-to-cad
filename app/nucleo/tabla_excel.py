"""Excel de la Tabla de datos, de IDA Y VUELTA (PURO: sin Qt).

`exportar` escribe una hoja por tabla (utilidades, buzones, curvas, bancoductos, avisos,
conflictos) más «Cotas por vértice» y una hoja «Léeme». Las columnas EDITABLES van en
verde; las grises son de consulta. `importar` lee el libro (encabezados en español o en
inglés), `cambios` compara con el proyecto (con la tolerancia del redondeo del Excel:
un valor que no se tocó no cuenta como cambio) y `aplicar` los pone en el modelo. La
ventana muestra los cambios antes de aplicarlos y se pueden deshacer.

Editables: utilidades (nombre, tipo, estado, familia, tamaño o diámetro, material, amperaje,
rasante inicial y final; familia y tamaño se validan contra el catálogo de Civil 3D), buzones y cajas
(tapa, fondo), curvas (radio; «Automático» = el radio por defecto) y cotas por vértice
(llegada y salida; vacío = automática, interpolada). Los conflictos se recalculan solos.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

from nucleo import avance, model_ops
from nucleo.tabla_datos import Tabla, _r
from traduccion.i18n_core import N_, TRANSLATIONS, t

EDITABLES = {
    "utilidades": {N_("Nombre"): ("name", "texto", 0), N_("Tipo"): ("tipo", "texto", 0),
                   N_("Estado"): ("ab", "estado", 0), N_("Familia"): ("pipe_family", "familia", 0),
                   N_("Tamaño"): ("pipe_size", "tamano", 0), N_("Diámetro (in)"): ("diam", "diametro", 0.0006),
                   N_("Material"): ("material", "texto", 0),
                   N_("Amperaje (A)"): ("amperaje", "num", 0.06),
                   N_("Rasante inicial (ft)"): ("inv_start", "num", 0.0006),
                   N_("Rasante final (ft)"): ("inv_end", "num", 0.0006),
                   N_("Pendiente (%)"): ("pendiente", "pendiente", 0.0015)},
    "buzones": {N_("Tapa (ft)"): ("rim", "num", 0.0006), N_("Fondo (ft)"): ("sump", "num", 0.0006)},
    "curvas": {N_("Radio (ft)"): ("radius_ft", "radio", 0.006)},
    "vertices": {N_("Cota de llegada (ft)"): ("in", "num", 0.0006),
                 N_("Cota de salida (ft)"): ("out", "num", 0.0006)},
}
TITULO_VERTICES = N_("Cotas por vértice")
VERDE, VERDE_CLARO, GRIS = "C6EFCE", "EBF7EE", "D9D9D9"


@dataclass
class Cambio:
    objeto: str            # "pipe" | "struct" | "vertice"
    indice: int            # utilidad o estructura
    campo: str             # clave del modelo
    antes: object
    despues: object
    etiqueta: str          # «Utilidad #3 · Rasante inicial (ft)»
    vertice: int = -1
    textos: tuple = None   # (antes, después) como se muestran (familia: su nombre, no su id)


# ─────────────────────────── exportar ───────────────────────────

def tabla_vertices(pipes, a_cad):
    """Cotas por vértice INTERIOR de cada utilidad (las puntas son la rasante inicial/final)."""
    tb = Tabla("vertices", TITULO_VERTICES, [
        N_("Utilidad N°"), N_("Nombre"), N_("Vértice"), N_("Cota de llegada (ft)"), N_("Cota de salida (ft)"),
        N_("Automática (ft)"), N_("X"), N_("Y")])
    for i, p in enumerate(pipes or []):
        if not i & 255 or i == len(pipes) - 1:
            avance.avisar("tabla_vertices", i + 1, len(pipes))        # barra de carga (Tabla de datos)
        pts = p.get("pts") or []
        if p.get("world") or len(pts) < 3:
            continue
        z0 = float(p.get("inv_start") or 0.0); z1 = float(p.get("inv_end") or z0)
        auto = model_ops.interp_vertex_z(pts, z0, z1, {})
        v_in = {int(k): v for k, v in (p.get("vertex_inv_in") or {}).items()}
        v_out = {int(k): v for k, v in (p.get("vertex_inv_out") or {}).items()}
        nombre = (p.get("name") or "").strip() or model_ops.nombre_por_defecto(p, i)
        for v in range(1, len(pts) - 1):
            x, y = a_cad(*pts[v])
            tb.filas.append([i + 1, nombre, v, _r(v_in.get(v), 3), _r(v_out.get(v), 3), _r(auto[v], 3),
                             _r(x, 3), _r(y, 3)])
            tb.lugares.append((pts[v][0], pts[v][1], "pipe", i))
    return tb


def exportar(tablas, ruta):
    """Un libro con una hoja por tabla; las columnas editables en verde."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ley = wb.active
    ley.title = t("Léeme")[:31]
    for fila in (
            [t("Cómo editar este Excel y volver a cargarlo")],
            [t("1. Cambia solo las columnas VERDES (cotas, tapas y fondos, radios, nombre, tipo, estado, familia, "
               "tamaño o diámetro, material, amperaje).")],
            ["   " + t("Familia y tamaño se escriben como en el catálogo de Civil 3D; un diámetro nuevo elige el "
                      "tamaño de esa medida en la familia de la utilidad. «Por defecto» = sin familia o tamaño.")],
            [t("2. Las columnas grises son de consulta: no se importan.")],
            [t("3. En la app: Herramientas → Tabla de datos → «Importar Excel…». Verás los cambios antes de "
               "aplicarlos y se pueden deshacer con Ctrl+Z.")],
            [t("4. En «Cotas por vértice», una celda vacía = cota automática (interpolada).")],
            [t("5. Pendiente (%) = (rasante final − rasante inicial) / largo × 100. Si la cambias, cambia también "
               "la rasante inicial o la final para que coincidan: si no coinciden, el Excel NO se importa.")],
            [t("6. Los conflictos (hoja «Conflictos», C-1, C-2…) se recalculan al importar.")]):
        ley.append(fila)
    ley["A1"].font = Font(bold=True, size=13)
    ley.column_dimensions["A"].width = 110
    # barra de carga (pedido 2026-10-20): filas escritas y celdas coloreadas, «i de n» de todo el libro
    total = sum(len(tb.filas) for tb in tablas) or 1
    celdas = sum(len(tb.filas) * len(EDITABLES.get(tb.clave, {})) for tb in tablas) or 1
    hechas, pintadas = 0, 0
    for tb in tablas:
        ws = wb.create_sheet(t(tb.titulo)[:31])
        editables = EDITABLES.get(tb.clave, {})
        ws.append([t(c) for c in tb.columnas])
        for f in tb.filas:
            ws.append(list(f))
            hechas += 1
            if not hechas & 255 or hechas == total:
                avance.avisar("excel_filas", hechas, total)
        for k, c in enumerate(tb.columnas, 1):
            cab = ws.cell(row=1, column=k)
            cab.font = Font(bold=True)
            cab.alignment = Alignment(wrap_text=True, vertical="center")
            editable = c in editables
            cab.fill = PatternFill("solid", fgColor=VERDE if editable else GRIS)
            if editable:
                for r in range(2, len(tb.filas) + 2):
                    ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=VERDE_CLARO)
                    pintadas += 1
                    if not pintadas & 1023 or pintadas == celdas:
                        avance.avisar("excel_formato", pintadas, celdas)
            largo = max([len(str(t(c)))] + [len(str(f[k - 1])) for f in tb.filas if f[k - 1] is not None])
            ws.column_dimensions[openpyxl.utils.get_column_letter(k)].width = min(60, max(9, largo + 2))
        ws.freeze_panes = "A2"
        if tb.filas:
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(tb.columnas))}{len(tb.filas) + 1}"
    with _contando_filas(total + len(tablas) + 8):
        wb.save(ruta)


def _contando_filas(total):
    """Mientras se guarda, avisa «excel_guardar» con cuántas filas escribió openpyxl (barra de carga; solo
    cuenta: el archivo es el mismo). Si esta versión de openpyxl no tiene `WorksheetWriter.write_row`, solo
    avisa que empieza."""
    import contextlib
    try:
        from openpyxl.worksheet._writer import WorksheetWriter
        original = WorksheetWriter.write_row
    except (ImportError, AttributeError):
        avance.avisar("excel_guardar")
        return contextlib.nullcontext()
    cuenta = [0]

    def write_row(self, xf, row, row_idx):
        cuenta[0] += 1
        if not cuenta[0] & 255:
            avance.avisar("excel_guardar", min(cuenta[0], total), total)
        return original(self, xf, row, row_idx)

    @contextlib.contextmanager
    def contando():
        avance.avisar("excel_guardar", 0, total)
        WorksheetWriter.write_row = write_row
        try:
            yield
        finally:
            WorksheetWriter.write_row = original
        avance.avisar("excel_guardar", total, total)
    return contando()


# ─────────────────────────── importar ───────────────────────────

def _norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _en_dos_idiomas(clave):
    en = TRANSLATIONS.get("en", {}).get(clave, clave)
    return {_norm(clave), _norm(en)}


def importar(ruta):
    """{clave de tabla: [ {columna (clave en español): valor} ]} de las hojas reconocidas."""
    import openpyxl
    from nucleo import tabla_datos
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    titulos = {"utilidades": N_("Utilidades"), "buzones": N_("Buzones y cajas"), "curvas": N_("Curvas"),
               "vertices": TITULO_VERTICES}
    columnas = {tb.clave: tb.columnas for tb in tabla_datos.armar([], [], [], [], lambda x, y: (x, y), 1.0)}
    columnas["vertices"] = tabla_vertices([], lambda x, y: (x, y)).columnas
    hojas = {_norm(n): n for n in wb.sheetnames}
    leido = {}
    for clave, titulo in titulos.items():
        nombre = next((hojas[n] for n in _en_dos_idiomas(titulo) if n in hojas), None)
        if nombre is None:
            continue
        filas = list(wb[nombre].iter_rows(values_only=True))
        if not filas:
            continue
        mapa = {}
        for k, cab in enumerate(filas[0]):
            for col in columnas[clave]:
                if _norm(cab) in _en_dos_idiomas(col):
                    mapa[k] = col
        leido[clave] = [{mapa[k]: v for k, v in enumerate(f) if k in mapa} for f in filas[1:]
                        if any(v not in (None, "") for v in f)]
    wb.close()
    return leido


def _num(v):
    """Número de una celda (acepta «-4,25»), None si vacía; ValueError si no es número."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).strip().replace(",", "."))


def _distinto(a, b, tol):
    if a is None or b is None:
        return (a is None) != (b is None)
    return abs(float(a) - float(b)) > tol


# ───────────── sección de la utilidad: estado, familia, tamaño, diámetro ─────────────

_SECCION = ("estado", "familia", "tamano", "diametro")
_A_BLOQUE = {"ab": "ab", "pipe_family": "familia", "pipe_size": "tamano"}
_POR_DEFECTO = {_norm(N_("Por defecto")), "default", ""}
_ACTIVA = {_norm(N_("Activa")), "active"}
_ABANDONADA = {_norm(N_("Abandonada")), "abandoned"}


def _txt(v):
    return "" if v is None else (f"{v:g}" if isinstance(v, float) else str(v)).strip()


def _cambios_seccion(p, i, fila, catalogo, con_bancoducto, out, errores):
    """Estado, familia y tamaño (o diámetro) de la utilidad i desde su fila del Excel."""
    from nucleo.edicion_bloque import diam_de_tamano
    etq = lambda col: t("Utilidad #{n} · {campo}").format(n=i + 1, campo=t(col))  # noqa: E731

    if N_("Estado") in fila and fila[N_("Estado")] not in (None, ""):
        v = _norm(fila[N_("Estado")])
        if v not in _ACTIVA | _ABANDONADA:
            errores.append(t("{donde}: «{v}» no es «Activa» ni «Abandonada».").format(
                donde=etq(N_("Estado")), v=fila[N_("Estado")]))
        elif (v in _ABANDONADA) != bool(p.get("ab")):
            out.append(Cambio("pipe", i, "ab", bool(p.get("ab")), v in _ABANDONADA, etq(N_("Estado")),
                              textos=(t("Abandonada") if p.get("ab") else t("Activa"),
                                      t("Abandonada") if v in _ABANDONADA else t("Activa"))))

    familias, tamanos_de = catalogo(p) if catalogo else ([], lambda _f: [])
    fam_actual = p.get("pipe_family") or ""
    fam = fam_actual
    nombres_actual = next((n for f, n in familias if f == fam_actual), set()) | {_norm(fam_actual)}
    if N_("Familia") in fila:
        v = _norm(fila[N_("Familia")])
        if (v in _POR_DEFECTO and not fam_actual) or (fam_actual and v in {_norm(n) for n in nombres_actual}):
            pass                                                       # sin tocar
        elif v in _POR_DEFECTO:
            fam = ""
        else:
            fam = next((f for f, n in familias if v in {_norm(x) for x in n}), None)
            if fam is None:
                errores.append(t("{donde}: «{v}» no está en el catálogo de Civil 3D de esta utilidad.").format(
                    donde=etq(N_("Familia")), v=fila[N_("Familia")]))
                fam = fam_actual

    from nucleo import model_ops as _mo
    tam_actual = p.get("pipe_size") or ""
    tam = tam_actual
    tamanos = tamanos_de(fam) if fam else []
    v = _txt(fila[N_("Tamaño")]) if N_("Tamaño") in fila else None
    editado = v is not None and not (_norm(v) == _norm(tam_actual) or (_norm(v) in _POR_DEFECTO and not tam_actual))
    d = None
    if N_("Diámetro (in)") in fila:
        try:
            d = _num(fila[N_("Diámetro (in)")])
        except ValueError:
            errores.append(t("{donde}: «{v}» no es un número.").format(donde=etq(N_("Diámetro (in)")),
                                                                       v=fila[N_("Diámetro (in)")]))
    d_actual = _mo.diametro(p)[0]
    d_cambia = d is not None and abs(d - d_actual) > 0.0006
    if editado:                                       # el tamaño escrito manda sobre el diámetro
        if _norm(v) in _POR_DEFECTO:
            tam = ""
        else:
            tam = next((x for x in tamanos if _norm(x) == _norm(v)), None)
            if tam is None and not tamanos:
                tam = v if re.match(r"\s*\d", v) else None
            if tam is None:
                errores.append(t("{donde}: «{v}» no es un tamaño de su familia.").format(
                    donde=etq(N_("Tamaño")), v=v))
                return
    elif d_cambia or fam != fam_actual:               # el diámetro elige el tamaño (en la familia nueva)
        meta = d if d is not None else d_actual
        if tamanos:
            tam = next((x for x in tamanos if abs(diam_de_tamano(x) - meta) <= 0.01), None)
        else:
            tam = f"diam:{meta:g}" if (d_cambia and not fam and meta > 0) else ("" if not d_cambia else None)
        if tam is None:
            if d_cambia:
                errores.append(t("{donde}: no hay un tamaño de {d} in en su familia.").format(
                    donde=etq(N_("Diámetro (in)")), d=f"{meta:g}"))
                return
            tam = ""                                  # familia nueva sin esa medida: su tamaño por defecto
    if (fam, tam) == (fam_actual, tam_actual):
        return
    if con_bancoducto:
        errores.append(t("Utilidad #{n}: lleva bancoducto; su sección la define el bancoducto.").format(n=i + 1))
        return
    nombre = lambda f: next((sorted(n, key=len)[0] for g, n in familias if g == f and n), f) or t("Por defecto")  # noqa: E731
    if fam != fam_actual:
        out.append(Cambio("pipe", i, "pipe_family", fam_actual, fam, etq(N_("Familia")),
                          textos=(nombre(fam_actual), nombre(fam))))
    if tam != tam_actual or fam != fam_actual:
        out.append(Cambio("pipe", i, "pipe_size", tam_actual, tam, etq(N_("Tamaño")),
                          textos=(tam_actual or t("Por defecto"),
                                  (f'{tam[5:]}"' if tam.startswith("diam:") else tam) or t("Por defecto"))))


def cambios(pipes, structures, leido, catalogo=None, con_bancoducto=()):
    """([Cambio], [errores]) entre el Excel leído y el proyecto. `catalogo(p)` →
    ([(id, {nombres})], tamanos_de(id) → [tamaño]) de las familias de Civil 3D para esa
    utilidad (None = sin catálogo: solo se aceptan «Por defecto» y diámetros sin familia)."""
    out, errores = [], []
    con_bancoducto = set(con_bancoducto or ())

    def valor(tipo, v, etiqueta):
        try:
            return _num(v) if tipo in ("num",) else v
        except ValueError:
            errores.append(t("{donde}: «{v}» no es un número.").format(donde=etiqueta, v=v))
            raise

    for fila in leido.get("utilidades", []):
        try:
            i = int(fila.get(N_("N°"))) - 1
        except (TypeError, ValueError):
            continue
        if not (0 <= i < len(pipes)):
            errores.append(t("Utilidad #{n}: no existe en el proyecto.").format(n=i + 1))
            continue
        p = pipes[i]
        _cambios_seccion(p, i, fila, catalogo, i in con_bancoducto, out, errores)
        for col, (campo, tipo, tol) in EDITABLES["utilidades"].items():
            if col not in fila or tipo in _SECCION or tipo == "pendiente":
                continue                                 # la pendiente se revisa en `incoherencias`
            etq = t("Utilidad #{n} · {campo}").format(n=i + 1, campo=t(col))
            try:
                nuevo = valor(tipo, fila[col], etq)
            except ValueError:
                continue
            viejo = p.get(campo)
            if tipo == "texto":
                nuevo = "" if nuevo is None else str(nuevo).strip()
                viejo = "" if viejo is None else str(viejo)
                if campo == "name" and nuevo == model_ops.nombre_por_defecto(p, i) and not viejo:
                    continue                             # el nombre por defecto no es un nombre propio
                if nuevo != viejo:
                    out.append(Cambio("pipe", i, campo, viejo, nuevo, etq))
            elif _distinto(viejo, nuevo, tol):
                out.append(Cambio("pipe", i, campo, viejo, nuevo, etq))

    por_cod = {s.get("cod"): n for n, s in enumerate(structures or []) if s.get("cod")}
    for clave in ("buzones", "curvas"):
        for fila in leido.get(clave, []):
            cod = fila.get(N_("Código"))
            if cod in (None, ""):
                continue
            n = por_cod.get(str(cod))
            if n is None:
                errores.append(t("«{cod}»: no existe en el proyecto.").format(cod=cod))
                continue
            s = structures[n]
            for col, (campo, tipo, tol) in EDITABLES[clave].items():
                if col not in fila:
                    continue
                etq = f"{cod} · {t(col)}"
                v = fila[col]
                if tipo == "radio":
                    try:
                        nuevo = _num(v)
                    except ValueError:
                        nuevo = 0.0                      # «Automático (…)»: el radio por defecto
                    nuevo = 0.0 if nuevo is None else nuevo
                    viejo = float(s.get("radius_ft") or 0.0)
                else:
                    try:
                        nuevo = valor(tipo, v, etq)
                    except ValueError:
                        continue
                    viejo = s.get(campo)
                if _distinto(viejo, nuevo, tol):
                    out.append(Cambio("struct", n, campo, viejo, nuevo, etq))

    for fila in leido.get("vertices", []):
        try:
            i = int(fila.get(N_("Utilidad N°"))) - 1
            v = int(fila.get(N_("Vértice")))
        except (TypeError, ValueError):
            continue
        if not (0 <= i < len(pipes)) or not (0 < v < len(pipes[i].get("pts") or []) - 1):
            errores.append(t("Utilidad #{n}, vértice {v}: no existe.").format(n=i + 1, v=v))
            continue
        p = pipes[i]
        for col, (lado, tipo, tol) in EDITABLES["vertices"].items():
            if col not in fila:
                continue
            etq = t("Utilidad #{n} · vértice {v} · {campo}").format(n=i + 1, v=v, campo=t(col))
            try:
                nuevo = valor(tipo, fila[col], etq)
            except ValueError:
                continue
            dic = {int(k): x for k, x in (p.get("vertex_inv_" + lado) or {}).items()}
            if _distinto(dic.get(v), nuevo, tol):
                out.append(Cambio("vertice", i, "vertex_inv_" + lado, dic.get(v), nuevo, etq, v))
    return out, errores


TOL_PENDIENTE = 0.0015       # % (el Excel la trae con 3 decimales)


def incoherencias(pipes, leido, ft_per_px):
    """Pendientes CAMBIADAS en el Excel que no coinciden con las rasantes de su fila (las
    nuevas o, si no se tocaron, las del proyecto). La pendiente se calcula con las cotas: si
    no coinciden, el Excel no se importa hasta que se corrija. Devuelve [mensaje]."""
    out = []
    for fila in leido.get("utilidades", []):
        if N_("Pendiente (%)") not in fila:
            continue
        try:
            i = int(fila.get(N_("N°"))) - 1
        except (TypeError, ValueError):
            continue
        if not (0 <= i < len(pipes)):
            continue
        p = pipes[i]
        try:
            pct = _num(fila[N_("Pendiente (%)")])
        except ValueError:
            out.append(t("Utilidad #{n}: la pendiente «{v}» no es un número.").format(
                n=i + 1, v=fila[N_("Pendiente (%)")]))
            continue
        actual = model_ops.pendiente_total(p, ft_per_px)
        if pct is None or actual is None or abs(pct - actual) <= TOL_PENDIENTE:
            continue                                     # no la cambiaron: se recalcula con las cotas
        L = model_ops.largo_planta_ft(p, ft_per_px)

        def rasante(col, campo):
            try:
                v = _num(fila.get(col))
            except ValueError:
                v = None
            return float(v) if v is not None else float(p.get(campo) or 0.0)
        zs, ze = rasante(N_("Rasante inicial (ft)"), "inv_start"), rasante(N_("Rasante final (ft)"), "inv_end")
        dan = (ze - zs) / L * 100.0
        if abs(pct - dan) <= TOL_PENDIENTE + 0.1 / L:     # + el redondeo de las rasantes a 3 decimales
            continue
        out.append(t("Utilidad #{n}: pendiente {p} %, pero sus rasantes ({zs} → {ze} ft en {L} ft) dan {dan} %. "
                     "Para {p} % la rasante final sería {sug} ft (o la inicial {sug0} ft).").format(
            n=i + 1, p=f"{pct:g}", zs=f"{zs:g}", ze=f"{ze:g}", L=f"{L:.2f}", dan=f"{dan:.3f}",
            sug=f"{zs + pct * L / 100.0:.3f}", sug0=f"{ze - pct * L / 100.0:.3f}"))
    return out


def aplicar(pipes, structures, lista):
    """Pone los cambios en el modelo (en sitio). Estado, familia y tamaño van por
    `edicion_bloque.aplicar` (el diámetro sale del tamaño, como en el panel)."""
    from nucleo import edicion_bloque
    seccion = {}
    for c in lista:
        if c.objeto == "pipe" and c.campo in _A_BLOQUE:
            seccion.setdefault(c.indice, {})[_A_BLOQUE[c.campo]] = c.despues
    for i, cb in seccion.items():
        if "familia" in cb and "tamano" not in cb:
            cb["tamano"] = ""
        edicion_bloque.aplicar(pipes, [i], cb)
    for c in lista:
        if c.objeto == "pipe" and c.campo in _A_BLOQUE:
            continue
        if c.objeto == "pipe":
            pipes[c.indice][c.campo] = c.despues
        elif c.objeto == "struct":
            structures[c.indice][c.campo] = c.despues
        elif c.objeto == "vertice":
            p = pipes[c.indice]
            dic = {int(k): x for k, x in (p.get(c.campo) or {}).items()}
            if c.despues is None:
                dic.pop(c.vertice, None)
            else:
                dic[c.vertice] = float(c.despues)
            p[c.campo] = dic


def texto(c):
    """«Utilidad #3 · Rasante inicial (ft): -4 → -5» (para la confirmación)."""
    def f(v):
        if v is None or v == "":
            return t("(vacío)")
        return f"{v:g}" if isinstance(v, float) else str(v)
    a, b = c.textos if c.textos else (c.antes, c.despues)
    return f"{c.etiqueta}: {f(a)} → {f(b)}"
