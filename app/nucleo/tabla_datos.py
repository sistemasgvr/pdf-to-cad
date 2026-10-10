"""Tabla con toda la información de cada elemento del proyecto (PURO: sin Qt).

Herramientas → «Tabla de datos…» (pedido del usuario 2026-10-09): utilidades, buzones y
cajas, curvas, bancoductos y avisos de normativa, cada uno con sus columnas, y
`exportar_excel` las escribe en un libro (una hoja por tabla, como tabla de Excel).

`armar(...)` recibe lo que solo la ventana sabe como funciones: `a_cad(x, y)` (px del
lienzo → coordenadas del DXF en pies, con la georreferencia si la hay),
`nombre_utilidad(capa)` y `nombre_familia(id)` (textos que se muestran).
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

from nucleo import model_ops
from nucleo.model import network_kind
from traduccion.i18n_core import N_, t


@dataclass
class Tabla:
    clave: str
    titulo: str                                   # N_(…)
    columnas: list                                # N_(…)
    filas: list = field(default_factory=list)
    lugares: list = field(default_factory=list)   # por fila: (x_px, y_px, objeto, índice) o None


RED = {"gravity": N_("Gravedad"), "conduit": N_("Conducto"), "pressure": N_("Presión")}


def _r(v, dec=2):
    try:
        return None if v is None or v == "" else round(float(v), dec)
    except (TypeError, ValueError):
        return v


def _xy(a_cad, pt):
    x, y = a_cad(pt[0], pt[1])
    return _r(x, 3), _r(y, 3)


def _largo_px(pts):
    return sum(math.dist(a, b) for a, b in zip(pts, pts[1:]))


def armar(pipes, structures, duct_banks, avisos, a_cad, ft_per_px,
          nombre_utilidad=lambda c: c, nombre_familia=lambda f: f):
    pipes = pipes or []
    structures = structures or []
    db_de = {}
    for db in duct_banks or []:
        for i in db.assigned():
            db_de[i] = db.name or t("(sin nombre)")
    avisos_de = {}
    for a in avisos or []:
        if not a.info:
            avisos_de[a.pipe] = avisos_de.get(a.pipe, 0) + 1

    ut = Tabla("utilidades", N_("Utilidades"), [
        N_("N°"), N_("Nombre"), N_("Utilidad"), N_("Tipo"), N_("Estado"), N_("Familia"), N_("Tamaño"),
        N_("Diámetro (in)"), N_("Largo (ft)"), N_("Vértices"), N_("Curvas"), N_("Rasante inicial (ft)"),
        N_("Rasante final (ft)"), N_("Amperaje (A)"), N_("Bancoducto"), N_("Inicio X"), N_("Inicio Y"),
        N_("Fin X"), N_("Fin Y"), N_("Avisos")])
    curva_de = {}
    for i, p in enumerate(pipes):
        pts = p.get("pts") or []
        if p.get("world") or len(pts) < 2:
            continue
        cv = model_ops.curve_vertex_indices(p, structures, pipes) if len(pts) >= 3 else {}
        for k, s in cv.items():
            curva_de[id(s)] = (i, k)
        d, defecto = model_ops.diametro(p)
        ix, iy = _xy(a_cad, pts[0])
        fx, fy = _xy(a_cad, pts[-1])
        ut.filas.append([
            i + 1, (p.get("name") or "").strip() or model_ops.nombre_por_defecto(p, i),
            nombre_utilidad(p.get("layer") or ""), p.get("tipo") or "",
            t("Abandonada") if p.get("ab") else t("Activa"),
            nombre_familia(p.get("pipe_family") or "") if p.get("pipe_family") else t("Por defecto"),
            p.get("pipe_size") or t("Por defecto"), _r(d, 3), _r(_largo_px(pts) * (ft_per_px or 0)), len(pts),
            len(cv), _r(p.get("inv_start"), 3), _r(p.get("inv_end"), 3), _r(p.get("amperaje"), 1),
            db_de.get(i, ""), ix, iy, fx, fy, avisos_de.get(i, 0)])
        ut.lugares.append((pts[len(pts) // 2][0], pts[len(pts) // 2][1], "pipe", i))

    bz = Tabla("buzones", N_("Buzones y cajas"), [
        N_("Código"), N_("Clase"), N_("Red"), N_("X"), N_("Y"), N_("Tapa (ft)"), N_("Fondo (ft)"),
        N_("Familia"), N_("Tamaño"), N_("Visible"), N_("Forma"), N_("Largo (ft)"), N_("Ancho (ft)"),
        N_("Alto (ft)")])
    cu = Tabla("curvas", N_("Curvas"), [
        N_("Código"), N_("Utilidad"), N_("Vértice"), N_("Radio (ft)"), N_("Origen"), N_("X"), N_("Y")])
    for n, s in enumerate(structures):
        if s.get("world") or s.get("x") is None:
            continue
        x, y = _xy(a_cad, (s["x"], s["y"]))
        if s.get("curve"):
            i, k = curva_de.get(id(s), (None, None))
            r = float(s.get("radius_ft") or 0.0)
            radio = _r(r) if r > 0.01 else (t("Automático ({r} ft)").format(
                r=_r(model_ops.radio_auto_ft(pipes[i]))) if i is not None else t("Automático"))
            cu.filas.append([s.get("cod", ""), (i + 1) if i is not None else "", k if k is not None else "",
                             radio, t("Quiebre del plano") if s.get("quiebre") else t("Curva"), x, y])
            cu.lugares.append((s["x"], s["y"], "curve", n))
            continue
        clase = t("Sólido") if s.get("solid") else (t("Caja") if s.get("net") == "conduit" else t("Buzón"))
        red = RED.get(s.get("net") or "gravity", "")
        bz.filas.append([
            s.get("cod", ""), clase, t(red) if red else "", x, y, _r(s.get("rim"), 3), _r(s.get("sump"), 3),
            s.get("part") or "", s.get("part_size") or "", t("No") if s.get("hidden") else t("Sí"),
            s.get("shape") or "", _r(s.get("length_ft")), _r(s.get("width_ft")),
            _r(s.get("solid_height_ft")) if s.get("solid") else None])
        bz.lugares.append((s["x"], s["y"], "struct", n))

    bd = Tabla("bancoductos", N_("Bancoductos"), [
        N_("Nombre"), N_("Ancho (in)"), N_("Alto (in)"), N_("Conductos"), N_("Utilidades")])
    for db in duct_banks or []:
        bd.filas.append([db.name or t("(sin nombre)"), _r(db.width_in), _r(db.height_in), len(db.conduits),
                         ", ".join(f"#{i + 1}" for i in db.assigned())])
        bd.lugares.append(None)

    av = Tabla("avisos", N_("Avisos de normativa"), [
        N_("N°"), N_("Utilidad"), N_("Tipo de aviso"), N_("Mensaje"), N_("X"), N_("Y")])
    from nucleo.normas_validar import CLASES
    for n, a in enumerate(avisos or [], 1):
        x, y = _xy(a_cad, (a.x, a.y))
        av.filas.append([n, (a.pipe + 1) if a.pipe is not None and a.pipe >= 0 else "",
                         t(CLASES.get(a.clase, a.clase)) + (" · " + t("Información") if a.info else ""),
                         a.mensaje, x, y])
        av.lugares.append((a.x, a.y, "pipe", a.pipe))
    return [ut, bz, cu, bd, av]


def exportar_excel(tablas, ruta):
    """Un libro con una hoja por tabla (encabezados ya traducidos)."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.worksheet.table import Table as TablaXL, TableStyleInfo
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for n, tb in enumerate(tablas, 1):
        ws = wb.create_sheet(t(tb.titulo)[:31])
        cab = [t(c) for c in tb.columnas]
        ws.append(cab)
        for f in tb.filas:
            ws.append(list(f))
        for c in ws[1]:
            c.font = Font(bold=True)
        for k, titulo in enumerate(cab, 1):
            largo = max([len(str(titulo))] + [len(str(f[k - 1])) for f in tb.filas if f[k - 1] is not None])
            ws.column_dimensions[openpyxl.utils.get_column_letter(k)].width = min(60, max(8, largo + 2))
        ws.freeze_panes = "A2"
        if tb.filas:
            ref = f"A1:{openpyxl.utils.get_column_letter(len(cab))}{len(tb.filas) + 1}"
            tabla = TablaXL(displayName=f"Tabla{n}", ref=ref)
            tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tabla)
    wb.save(ruta)
