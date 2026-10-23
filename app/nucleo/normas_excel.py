"""Excel de normativas (PURO: sin Qt; usa openpyxl). Desde 2026-10-21 = la tabla única de reglas
(`normas_reglas`): hojas INSTRUCCIONES, TIPOS y REGLAS (+ LISTAS, oculta, con lo que alimenta los desplegables).

Para que el ingeniero siga cargando datos SIN poder escribir lo que la app no contempla: cada columna que elige
de una lista tiene su desplegable con error que no deja seguir (UTILIDAD; TIPO y CON TIPO = los de la hoja
TIPOS para la utilidad de esa fila; ELEMENTO = los que existen en esa utilidad; REVISAR = lo que la app mide de
ese elemento en esa utilidad; NIVEL; ACTIVA; CON UTILIDAD), y las de números solo aceptan números. VALORES y SI
MATERIAL aceptan listas con «;» (se revisan al importar). Al importar, cada fila pasa por `normas_reglas.leer_fila`
—las MISMAS reglas que el editor de la ventana— y lo que no se entiende se anota (hoja, fila, columna) y no se usa.

Un Excel VIEJO de cuatro hojas (TIPOS, PRESION-DIAMETROS, PRESION-ACCESORIOS, ELECT-TELECOM) se sigue
importando: `normas_excel_v1` lo lee y `normas_migrar` lo pasa a reglas.
"""
from __future__ import annotations

from nucleo import normas_catalogo as nc
from nucleo import normas_excel_ayuda as ayuda
from nucleo import normas_reglas as NR

HOJA_TIPOS = "TIPOS"
HOJA_REGLAS = "REGLAS"
HOJA_LISTAS = "LISTAS"
COL_TIPOS = ("UTILIDAD", "TIPO")
AUX = [chr(ord("K") + k) for k in range(len(nc.UTILIDADES))]   # bloque auxiliar de TIPOS (una por utilidad)
FILAS_DV = 1000                     # filas con desplegable (la tabla crece dentro)
ESTILO_TABLA = "TableStyleMedium10"
_FILA_REVISAR = 40                  # en LISTAS: encabezados «UTILIDAD|ELEMENTO» de las listas de REVISAR


def _letra(k):
    from openpyxl.utils import get_column_letter
    return get_column_letter(k + 1)


def _nombre_lista(utilidad, elemento=None):
    """Nombre definido de Excel para una lista («E_AGUA», «R_AGUA_TDOMICILIARIA»)."""
    if elemento is None:
        return f"E_{utilidad}"
    return f"R_{utilidad}_" + elemento.replace(" ", "").replace("-", "")


# ─────────────────────────── importar ───────────────────────────

def _hoja(wb, nombre):
    for ws in wb.worksheets:
        if nc.clave(ws.title) == nc.clave(nombre):
            return ws
    return None


def importar(ruta, problemas=None):
    """(catálogo, hojas que faltan). Con la lista `problemas`, se le agrega cada celda que no se entendió (y cuya
    fila por eso no se usa)."""
    import openpyxl
    problemas = problemas if problemas is not None else []
    wb = openpyxl.load_workbook(ruta, data_only=True)
    if _hoja(wb, HOJA_REGLAS) is None:
        from nucleo import normas_excel_v1, normas_migrar
        viejo, faltan = normas_excel_v1.importar(ruta, problemas)
        if len(faltan) < 4:                            # es un Excel viejo: se pasa a reglas
            return nc.normalizar(normas_migrar.de_v1(viejo)), []
        return nc.normalizar({"tipos": [], "reglas": []}), [HOJA_TIPOS, HOJA_REGLAS]
    faltan = [] if _hoja(wb, HOJA_TIPOS) is not None else [HOJA_TIPOS]
    tipos, filas_t = [], _filas(_hoja(wb, HOJA_TIPOS), COL_TIPOS, problemas, conocidas=nc.UTILIDADES)
    for n, d in filas_t:
        u, t = nc.clave(d.get("UTILIDAD")), nc.clave(d.get("TIPO"))
        if u not in nc.UTILIDADES:
            problemas.append(f"{HOJA_TIPOS}, fila {n}, UTILIDAD: «{d.get('UTILIDAD')}» no es una utilidad de la "
                             f"aplicación")
        elif not t:
            problemas.append(f"{HOJA_TIPOS}, fila {n}, TIPO: está vacío")
        elif [u, t] not in tipos:
            tipos.append([u, t])
    estrictos = {tuple(x) for x in tipos} if not faltan else None
    reglas = []
    for n, d in _filas(_hoja(wb, HOJA_REGLAS), NR.COLUMNAS, problemas):
        r, mal = NR.leer_fila(d, estrictos)
        for col, motivo in mal:
            problemas.append(f"{HOJA_REGLAS}, fila {n}, {col}: {motivo}")
        if r is not None:
            reglas.append(r)
    return nc.normalizar({"tipos": tipos, "reglas": reglas}), faltan


def _filas(ws, columnas, problemas, conocidas=()):
    """[(nº de fila, {columna: valor})] con los encabezados de la fila 1; avisa las columnas desconocidas."""
    if ws is None:
        return []
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    cab = [nc.clave(c) if c is not None else "" for c in filas[0]]
    validas = {nc.clave(c) for c in columnas} | {nc.clave(c) for c in conocidas}
    for v in filas[0]:
        if v not in (None, "") and nc.clave(v) not in validas and not str(v).strip().startswith("_"):
            problemas.append(f"{ws.title}: columna «{v}» desconocida")
    por = {nc.clave(c): c for c in columnas}
    out = []
    for n, f in enumerate(filas[1:], 2):
        d = {por[cab[i]]: v for i, v in enumerate(f) if i < len(cab) and cab[i] in por}
        if any(v not in (None, "") and str(v).strip() for v in d.values()):
            out.append((n, d))
    return out


# ─────────────────────────── exportar ───────────────────────────

def _celda(col, v):
    if v in (None, ""):
        return None
    if col in NR.NUMERICAS:
        n = nc.numero(v)
        return int(n) if n is not None and float(n).is_integer() else n
    if col == "VALORES":
        nums = nc.lista_numeros(v)
        if len(nums) == 1 and ";" not in str(v) and nc.numero(v) is not None:
            return int(nums[0]) if float(nums[0]).is_integer() else nums[0]
    return v


def _tabla(ws, nombre, cabecera, filas, anchos):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.table import Table, TableStyleInfo
    ws.append(list(cabecera))
    for k, col in enumerate(cabecera):
        texto = ayuda.comentario(ws.title, col)
        if texto:
            com = Comment(texto, "pdf-to-cad")
            com.width, com.height = 340, 40 + 16 * (texto.count("\n") + len(texto) // 48)
            ws.cell(row=1, column=k + 1).comment = com
    for f in filas:
        ws.append([_celda(c, v) for c, v in zip(cabecera, f)])
    for i, w in enumerate(anchos):
        ws.column_dimensions[_letra(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    t = Table(displayName=nombre, ref=f"A1:{_letra(len(cabecera) - 1)}{max(2, len(filas) + 1)}")
    t.tableStyleInfo = TableStyleInfo(name=ESTILO_TABLA, showRowStripes=True)
    ws.add_table(t)


def _validacion(ws, col, tipo, formula, titulo, error, aviso=False, ayuda_txt=None, columnas=NR.COLUMNAS):
    from openpyxl.worksheet.datavalidation import DataValidation
    k = columnas.index(col)
    letra = _letra(k)
    if tipo == "decimal":
        dv = DataValidation(type="decimal", operator="between", formula1="-1000000", formula2="1000000")
    elif tipo == "list":
        dv = DataValidation(type="list", formula1=formula)
    else:
        dv = DataValidation(type="custom", formula1=formula)
    dv.allow_blank = True
    dv.showErrorMessage, dv.errorTitle, dv.error = True, titulo, error
    dv.errorStyle = "warning" if aviso else "stop"
    if ayuda_txt:
        dv.showInputMessage, dv.promptTitle, dv.prompt = True, titulo[:32], ayuda_txt[:255]
    dv.add(f"{letra}2:{letra}{FILAS_DV}")
    ws.add_data_validation(dv)


def _listas(wb):
    """Hoja oculta LISTAS: los elementos de cada utilidad y lo que se revisa de cada (utilidad, elemento), con
    nombres definidos (E_…, R_…) para los desplegables que dependen de la fila."""
    from openpyxl.workbook.defined_name import DefinedName
    wl = wb.create_sheet(HOJA_LISTAS)
    col = 0
    for u in nc.UTILIDADES:
        els = NR.elementos_de(u)
        wl.cell(row=1, column=col + 1, value=u)
        for n, e in enumerate(els, 2):
            wl.cell(row=n, column=col + 1, value=e)
        L = _letra(col)
        wb.defined_names[_nombre_lista(u)] = DefinedName(
            _nombre_lista(u), attr_text=f"{HOJA_LISTAS}!${L}$2:${L}${len(els) + 1}")
        col += 1
    col = 0
    for u in nc.UTILIDADES:
        for e in NR.elementos_de(u):
            rev = NR.revisar_de(u, e)
            wl.cell(row=_FILA_REVISAR, column=col + 1, value=f"{u}|{e}")
            for n, r in enumerate(rev, _FILA_REVISAR + 1):
                wl.cell(row=n, column=col + 1, value=r.texto)
            L = _letra(col)
            wb.defined_names[_nombre_lista(u, e)] = DefinedName(
                _nombre_lista(u, e), attr_text=f"{HOJA_LISTAS}!${L}${_FILA_REVISAR + 1}:${L}${_FILA_REVISAR + len(rev)}")
            col += 1
    wl.sheet_state = "hidden"
    return wl


def exportar(cat, ruta):
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    cat = nc.normalizar(cat)
    tipos = sorted((list(t) for t in cat.get("tipos", [])),
                   key=lambda t: nc.UTILIDADES.index(t[0]) if t[0] in nc.UTILIDADES else 99)
    wb = openpyxl.Workbook()
    wt = wb.active
    wt.title = HOJA_TIPOS
    _tabla(wt, "TablaTipos", COL_TIPOS, tipos, [18, 34])
    for col, u in zip(AUX, nc.UTILIDADES):             # bloque auxiliar oculto: los tipos de cada utilidad
        wt[f"{col}1"] = u
        wt[f"{col}2"] = ArrayFormula(f"{col}2:{col}201", f'=_xlfn._xlws.SORT(_xlfn.UNIQUE(_xlfn._xlws.FILTER('
                                                         f'TablaTipos[TIPO],TablaTipos[UTILIDAD]={col}$1,"")))')
        wt.column_dimensions[col].hidden = True
    wr = wb.create_sheet(HOJA_REGLAS)
    anchos = [8, 15, 26, 16, 9, 9, 9, 9, 9, 9, 16, 40, 9, 9, 24, 16, 22, 9, 30, 50]
    _tabla(wr, "TablaReglas", NR.COLUMNAS, [NR.fila_de(r) for r in cat.get("reglas", [])], anchos)
    _listas(wb)

    util = '"' + ",".join(nc.UTILIDADES) + '"'
    _validacion(wt, "UTILIDAD", "list", util, "Utilidad", "Elige una utilidad de la lista.", columnas=COL_TIPOS)

    def tipo_de(col):
        c = _letra(NR.COLUMNAS.index(col))
        m = f"MATCH(${c}2,TIPOS!${AUX[0]}$1:${AUX[-1]}$1,0)-1"
        return f'OFFSET(TIPOS!${AUX[0]}$2,0,{m},MAX(1,COUNTIF(OFFSET(TIPOS!${AUX[0]}$2,0,{m},200,1),"?*")),1)'

    U, E = _letra(NR.COLUMNAS.index("UTILIDAD")), _letra(NR.COLUMNAS.index("ELEMENTO"))
    no_lista = "Elige un valor de la lista: es lo que la aplicación sabe revisar."
    _validacion(wr, "ACTIVA", "list", '"SI,NO"', "Activa", "Escribe SI o NO.")
    _validacion(wr, "UTILIDAD", "list", util, "Utilidad", "Elige una utilidad de la lista.")
    _validacion(wr, "TIPO", "list", tipo_de("UTILIDAD"), "Tipo",
                "Elige un tipo de la lista (los tipos se crean en la hoja TIPOS) o déjala vacía.")
    _validacion(wr, "ELEMENTO", "list", f'INDIRECT("E_"&${U}2)', "Elemento", no_lista)
    _validacion(wr, "REVISAR", "list",
                f'INDIRECT("R_"&${U}2&"_"&SUBSTITUTE(SUBSTITUTE(${E}2," ",""),"-",""))', "Revisar",
                no_lista + " Primero elige UTILIDAD y ELEMENTO.")
    for col in sorted(NR.NUMERICAS, key=NR.COLUMNAS.index):
        _validacion(wr, col, "decimal", None, col[:32], "Escribe UN número, sin unidades ni símbolos. Para un rango "
                                                          "usa las dos columnas MIN y MAX.")
    _validacion(wr, "SI MATERIAL", "list", '"' + ",".join(NR.MATERIALES) + '"', "Material",
                "Materiales de la app. Para varios, sepáralos con «;».", aviso=True)
    _validacion(wr, "VALORES", "custom", "TRUE", "Valores", "", ayuda_txt=ayuda.COMENTARIOS["VALORES"])
    _validacion(wr, "CON UTILIDAD", "list", '"' + ",".join(nc.UTILIDADES + (NR.CUALQUIERA,)) + '"',
                "Con utilidad", "Elige la otra utilidad de la lista (o CUALQUIERA).")
    _validacion(wr, "CON TIPO", "list", tipo_de("CON UTILIDAD"), "Con tipo",
                "Elige un tipo de la otra utilidad o déjala vacía.")
    _validacion(wr, "NIVEL", "list", '"' + ",".join(NR.NIVELES) + '"', "Nivel", "Escribe ERROR, AVISO o INFO.")
    ayuda.escribir_instrucciones(wb)
    wb.active = 0
    wb.save(ruta)
