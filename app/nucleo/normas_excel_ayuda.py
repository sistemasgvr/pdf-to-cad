"""Textos de ayuda del Excel de normativas (PURO): el comentario de cada columna y la hoja INSTRUCCIONES.
Desde 2026-10-21 el Excel es UNA hoja de reglas (REGLAS) + TIPOS (`normas_reglas`); lo que la app sabe medir
está en una lista cerrada y las celdas que eligen de una lista no dejan escribir otra cosa.

La hoja INSTRUCCIONES (y cualquier otra que no sea TIPOS o REGLAS, como «auxiliar») se escribe al exportar y
NUNCA se lee al importar. Los textos van en español fijo: el Excel es un formato de datos (sus encabezados
también son fijos).
"""
from __future__ import annotations

from nucleo import normas_reglas as NR

HOJA_INSTRUCCIONES = "INSTRUCCIONES"

_VACIO = "Vacía = sin restricción."
_MIN = "Un solo número, sin unidades. Vacía = sin mínimo. El valor escrito también vale (es «desde»)."
_MAX = "Un solo número, sin unidades. Vacía = sin máximo. El valor escrito también vale (es «hasta»)."
_SI = "Condición: la regla solo vale para las tuberías que la cumplen. " + _VACIO

COMENTARIOS = {
    "UTILIDAD": "Elige de la lista: AGUA, ALCANTARILLADO, DRENAJE, GAS, ELECTRICO o TELECOM.",
    "TIPO": "Elige de la lista (los tipos se crean en la hoja TIPOS). Vacía = todos los tipos de la utilidad. "
            "En un accesorio, es el tipo de su tubería PRINCIPAL.",
    "ACTIVA": "SI = se revisa. NO = queda escrita pero no se revisa.",
    "ELEMENTO": "Qué se mira: TUBERIA, CURVA, BUZON (redes por gravedad) o un accesorio (redes a presión): "
                "Y, T, CRUZ, CODO, T - DOMICILIARIA (la T cuyo ramal es de instalación domiciliaria).",
    "SI DIAMETRO MIN (in)": _SI + "\nDiámetro de la tubería desde (pulgadas).",
    "SI DIAMETRO MAX (in)": _SI + "\nDiámetro de la tubería hasta (pulgadas).",
    "SI AMPERAJE MIN (A)": _SI + "\nAmperaje desde (solo eléctricas). Ejemplo «de 0 a 320»: MIN 0, MAX 320.",
    "SI AMPERAJE MAX (A)": _SI + "\nAmperaje hasta (solo eléctricas).",
    "SI LARGO MIN (ft)": _SI + "\nLargo de la tubería en planta desde (pies).",
    "SI LARGO MAX (ft)": _SI + "\nLargo de la tubería en planta hasta (pies).",
    "SI MATERIAL": _SI + "\nMaterial(es) de la app separados con «;»: " + "; ".join(NR.MATERIALES) + ".",
    "REVISAR": "Elige de la lista lo que se mide (cambia según UTILIDAD y ELEMENTO). La unidad va en el nombre. "
               "«Solo referencia» = queda escrita con su fuente, sin revisar.",
    "MIN": "Lo medido debe ser al menos esto (unidad de REVISAR).\n" + _MIN,
    "MAX": "Lo medido no debe pasar de esto (unidad de REVISAR).\n" + _MAX,
    "VALORES": "Lo medido debe ser UNO de estos, separados con «;» (6; 8; 12 o 11,25; 22,5; 45). Según REVISAR: "
               "números, tipos de la utilidad, materiales de la app, o SI / NO en «Se permite».",
    "CON UTILIDAD": "Solo en separaciones: respecto a qué otra utilidad. CUALQUIERA = cualquier otra.",
    "CON TIPO": "Solo en separaciones: el tipo de la otra utilidad. " + _VACIO,
    "NIVEL": "ERROR, AVISO o INFO (INFO solo informa, no es un incumplimiento). Vacía = AVISO.",
    "FUENTE": "De dónde sale la regla (norma, manual, página). Texto libre.",
    "NOTA": "Aclaración para quien revisa. En un Diámetro, la NOTA dice dónde se usa: si un diámetro solo lo "
            "permiten reglas con NOTA (4 in en cul-de-sacs), la app la muestra como información.",
}


def comentario(hoja, columna):
    return COMENTARIOS.get(columna, "")


# Qué mide cada opción de REVISAR, en palabras sencillas (una por clave de `normas_reglas.REVISAR`; una prueba
# exige que no falte ninguna).
EXPLICA = {
    "diametro": "El diámetro de la tubería.",
    "pendiente": "La pendiente de cada tramo: cuántos pies baja (o sube) por cada 100 pies de largo. Necesita las "
                 "cotas.",
    "largo": "El largo de la tubería visto en planta.",
    "amperaje": "El amperaje escrito en la línea eléctrica.",
    "material": "El material debe ser uno de los de VALORES.",
    "material_no": "El material NO puede ser ninguno de los de VALORES.",
    "quiebre": "El giro más grande que puede tener la tubería en un vértice sin buzón ni curva (en MAX).",
    "llega_menos": "Cuando esta tubería termina sobre otra de la misma utilidad: cuántas pulgadas más chica es que "
                   "esa otra.",
    "sep_h": "Distancia entre dos tuberías que van una al lado de la otra, medida de pared a pared (descuenta el "
             "grosor de cada una). Se pone en MIN y con CON UTILIDAD.",
    "sep_h_ejes": "Lo mismo, pero medida de centro a centro de cada tubería.",
    "sep_v": "Donde las dos tuberías se cruzan: el espacio libre de alto entre ellas, vaya cual vaya arriba. "
             "Necesita las cotas.",
    "encima": "Donde se cruzan: esta tubería debe pasar POR ENCIMA de la otra, con al menos MIN pies libres.",
    "debajo": "Donde se cruzan: esta tubería debe pasar POR DEBAJO de la otra, con al menos MIN pies libres.",
    "radio": "El radio de la curva. Si la curva no tiene radio escrito, se mira el automático.",
    "dist_buzones": "El largo de tubería entre dos buzones seguidos.",
    "caida": "En el buzón, la diferencia entre la cota con que llega la tubería y la cota con que sale.",
    "diam_buzon": "Diámetro nominal del buzón: su medida reconocida del plano o el tamaño elegido (sin dato, no se revisa).",
    "permitido": "Si el accesorio se puede usar: SI o NO (en VALORES).",
    "solo_en": "El accesorio solo se puede usar en esos tipos de tubería principal (en VALORES).",
    "diam_principal": "El diámetro de la tubería principal del accesorio.",
    "diam_ramal": "El diámetro del ramal (la tubería que sale de lado).",
    "ramal_pasos": "El ramal debe ser un tamaño más chico que la principal: escribe -1 en VALORES (se cuenta en "
                   "la lista de «Diámetro de la principal»).",
    "ramal_menos": "Cuántas pulgadas más chico que la principal es el ramal.",
    "angulo": "En el CODO, cuánto gira la tubería (90 = vuelta en escuadra). En la T, la Y y la CRUZ, el ángulo "
              "entre el ramal y la principal.",
    "referencia": "No se mide: la regla queda escrita con su FUENTE y su NOTA para consulta.",
}

_COLUMNAS_SE = {"n": "MIN", "x": "MAX", "v": "VALORES"}
_NOMBRE = {"ELECTRICO": "eléctricas", "TELECOM": "telecom"}


def _tabla_revisar():
    filas = [(["Qué revisar (REVISAR)", "Para qué ELEMENTO", "Se escribe en", "Qué significa"], "h")]
    for r in NR.REVISAR:
        para = "Accesorios" if set(r.elementos) == set(NR.ACCESORIOS) else \
            "Todos" if set(r.elementos) == set(NR.ELEMENTOS) else ", ".join(r.elementos)
        if r.redes == ("gravity",):
            para += " (alcantarillado y drenaje)"
        if r.utilidades:
            para += " (solo " + ", ".join(_NOMBRE.get(u, u.lower()) for u in r.utilidades) + ")"
        se = " / ".join(_COLUMNAS_SE[c] for c in r.usa) or "—"
        filas.append(([r.texto, para, se, EXPLICA[r.clave]], "f"))
    return filas


# (texto, estilo): "t" título, "s" subtítulo, "" normal, "e" ejemplo, "h"/"f" encabezado/fila de una tabla.
INSTRUCCIONES = [
    ("Cómo llenar este Excel de normativas", "t"),
    ("Este Excel guarda las reglas de diseño que la app revisa en el plano (diámetros, pendientes, separaciones, "
     "accesorios…). Cuando algo no cumple una regla, la app lo marca en el plano con un aviso.", ""),
    ("Las mismas reglas también se pueden ver y cambiar dentro de la app, sin Excel: menú Normativas → "
     "Normativas de diseño (pestañas Tipos y Reglas). El Excel sirve para revisarlas o llenarlas de a muchas.", ""),
    ("", ""),
    ("1. Las hojas", "s"),
    ("TIPOS: los tipos de tubería de cada utilidad (por ejemplo, en AGUA: Distribución principal, Instalación "
     "domiciliaria…).", ""),
    ("REGLAS: una fila = una regla.", ""),
    ("Esta hoja (INSTRUCCIONES) es solo ayuda: la app no la lee. Hay además una hoja oculta (LISTAS) con lo que "
     "alimenta las listas desplegables: no la cambies.", ""),
    ("", ""),
    ("2. Cómo se lee una fila de REGLAS", "s"),
    ("«En esta UTILIDAD y TIPO, para este ELEMENTO, si se cumplen las condiciones SI…, lo que dice REVISAR debe "
     "estar entre MIN y MAX, o ser uno de VALORES.»", ""),
    ("Ejemplos:", ""),
    ("ALCANTARILLADO · COLECTOR · TUBERIA · SI DIAMETRO 8 a 8 · Pendiente (%) · MIN 0,44", "e"),
    ("   → En los colectores de alcantarillado de 8\", la pendiente debe ser de al menos 0,44 %.", "e"),
    ("AGUA · DISTRIBUCION PRINCIPAL · CODO · Ángulo (°) · VALORES 11,25; 22,5; 45; 90", "e"),
    ("   → En la distribución principal de agua, los codos solo pueden ser de 11,25°, 22,5°, 45° o 90°.", "e"),
    ("AGUA · TUBERIA · Separación horizontal en paralelo, entre bordes (ft) · MIN 10 · CON ALCANTARILLADO", "e"),
    ("   → Una tubería de agua que va al lado de una de alcantarillado debe quedar al menos a 10 pies de ella.", "e"),
    ("", ""),
    ("3. Qué va en cada columna", "s"),
    ("Pasa el mouse por el título de cada columna en la hoja REGLAS para ver esta misma ayuda.", ""),
    ("ACTIVA — SI = la app revisa esta regla. NO = la regla queda escrita pero no se revisa (sirve para apagarla "
     "un tiempo sin borrarla).", ""),
    ("UTILIDAD — A qué red se aplica: AGUA, ALCANTARILLADO, DRENAJE, GAS, ELECTRICO o TELECOM.", ""),
    ("TIPO — A qué tipo de tubería de esa utilidad (los de la hoja TIPOS). Vacío = a todos los tipos. En un "
     "accesorio es el tipo de su tubería principal.", ""),
    ("ELEMENTO — Qué se mira: la TUBERIA, una CURVA, un BUZON o un accesorio (ver el punto 4).", ""),
    ("SI … (DIAMETRO, AMPERAJE, LARGO, MATERIAL) — Condiciones: la regla solo vale para las tuberías que las "
     "cumplen. Vacías = vale para todas. Ejemplo: SI DIAMETRO MIN 8 y SI DIAMETRO MAX 12 = solo tuberías de 8\" "
     "a 12\".", ""),
    ("REVISAR — Qué se mide (ver la tabla del punto 5). Se elige de la lista; la lista cambia según la UTILIDAD y "
     "el ELEMENTO de la fila.", ""),
    ("MIN — Lo medido debe ser al menos esto.  MAX — Lo medido no debe pasar de esto.  VALORES — Lo medido debe "
     "ser uno de estos. Puedes llenar MIN, MAX, VALORES o varios a la vez.", ""),
    ("CON UTILIDAD y CON TIPO — Solo en las separaciones: respecto a qué otra utilidad (y, si quieres, a qué "
     "tipo de ella). CUALQUIERA = cualquier otra utilidad. En las demás reglas van vacías.", ""),
    ("NIVEL — Qué tan grave es no cumplir: ERROR = grave (aviso en rojo); AVISO = normal (aviso ámbar); INFO = "
     "no es un incumplimiento, solo informa (aviso celeste con «i»). Vacío = AVISO.", ""),
    ("FUENTE — De dónde sale la regla (norma, manual, página). Texto libre; aparece en la lista de avisos y en el informe.",
     ""),
    ("NOTA — Aclaración para quien revisa. En un diámetro, la NOTA dice dónde se usa: si un diámetro solo lo "
     "permiten reglas con NOTA (por ejemplo 4\" en cul-de-sacs), la app no lo marca como incumplimiento, solo "
     "muestra la nota como información.", ""),
    ("", ""),
    ("4. Qué ELEMENTO tiene cada utilidad", "s"),
    ("TUBERIA y CURVA: en todas las utilidades.", ""),
    ("BUZON: solo en las redes por gravedad (ALCANTARILLADO y DRENAJE).", ""),
    ("Accesorios — Y, T, CRUZ, CODO y T - DOMICILIARIA (la T cuyo ramal es de instalación domiciliaria): solo en "
     "las redes a presión (AGUA y GAS).", ""),
    ("ELECTRICO y TELECOM: solo TUBERIA y CURVA.", ""),
    ("", ""),
    ("5. Qué se puede revisar", "s"),
    ("La app solo acepta en REVISAR lo de esta tabla: es lo que sabe medir en el plano.", ""),
    *_tabla_revisar(),
    ("", ""),
    ("6. Varias reglas de lo mismo", "s"),
    ("Si dos filas son iguales en UTILIDAD, TIPO, ELEMENTO, condiciones SI…, REVISAR y CON…, basta con cumplir "
     "UNA de ellas. Ejemplo: una fila con diámetros 6; 8; 12; 16 y otra con 4 (para cul-de-sacs): un diámetro de "
     "4 cumple. Si algo cambia entre las filas, se revisan las dos.", ""),
    ("", ""),
    ("7. Cómo escribir números y listas", "s"),
    ("Solo el número, sin unidades ni símbolos: 4   (no 4\", no 4 in). Decimales con coma o punto: 1,5", "e"),
    ("Listas con punto y coma: 6; 8; 12; 16", "e"),
    ("Rangos en dos columnas: de 0,5 a 2 → MIN 0,5 y MAX 2. «4 o más» → MIN 4 y MAX vacía.", "e"),
    ("El número escrito también vale: MIN 4 acepta el 4; MAX 12 acepta el 12.", "e"),
    ("", ""),
    ("8. Tipos", "s"),
    ("Para crear un tipo nuevo, agrega una fila en la hoja TIPOS con su UTILIDAD y su TIPO. Después ese tipo "
     "aparece en las listas de REGLAS (también se pueden crear dentro de la app).", ""),
    ("", ""),
    ("9. Lo que la app no puede medir", "s"),
    ("Velocidades, caudales, presiones, profundidades sin terreno, camas o sondeos: se escriben con REVISAR «Solo "
     "referencia» (quedan escritos con su FUENTE y no dan avisos).", ""),
    ("Pendientes, caídas en el buzón y cruces se miden con las cotas de cada tubería: si una tubería no tiene "
     "cotas, esas reglas no se revisan en ella.", ""),
    ("", ""),
    ("10. Al importar", "s"),
    ("En la app: Normativas → Normativas de diseño → «Importar Excel…». Solo se leen las hojas TIPOS y REGLAS.",
     ""),
    ("No cambies los títulos de las columnas: la app las busca por su nombre.", ""),
    ("Si una celda no se entiende o no es algo que la app contempla, la app te dice cuál (hoja, fila y columna) y "
     "esa fila no se usa (las demás sí).", ""),
    ("Un Excel viejo de cuatro hojas (PRESION-DIAMETROS, PRESION-ACCESORIOS…) también se puede importar: se pasa "
     "solo a reglas.", ""),
]

_ANCHOS = (44, 30, 18, 90)


def escribir_instrucciones(wb):
    """Agrega la hoja INSTRUCCIONES como PRIMERA hoja del libro."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    ws = wb.create_sheet(HOJA_INSTRUCCIONES, 0)
    for k, ancho in enumerate(_ANCHOS):
        ws.column_dimensions["ABCD"[k]].width = ancho
    ws.sheet_view.showGridLines = False
    borde = Border(*(Side(style="thin", color="B7C9DD"),) * 4)
    for texto, estilo in INSTRUCCIONES:
        if estilo in ("h", "f"):
            ws.append(list(texto))
            for k in range(len(texto)):
                c = ws.cell(row=ws.max_row, column=k + 1)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = borde
                if estilo == "h":
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="1F5C99")
            continue
        ws.append([texto])
        fila = ws.max_row
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(_ANCHOS))
        c = ws.cell(row=fila, column=1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        largo = len(texto)
        if estilo == "t":
            c.font = Font(bold=True, size=15)
        elif estilo == "s":
            c.font = Font(bold=True, size=12, color="1F5C99")
        elif estilo == "e":
            c.font = Font(name="Consolas", size=11)
            c.fill = PatternFill("solid", fgColor="EEF4FB")
            largo = int(largo * 1.25)
        # Excel no agranda solo una fila de celdas combinadas: alto según el largo del texto.
        ws.row_dimensions[fila].height = 15 * max(1, -(-largo // 165))
    return ws
